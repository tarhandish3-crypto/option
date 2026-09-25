# 0myStrategy/covered_call_analyzer.py
# -*- coding: utf-8 -*-


import sys
from pathlib import Path

current_file_path = Path(__file__).resolve()
current_dir = current_file_path.parent
root_dir = current_dir.parent
sys.path.append(str(root_dir))

import json
import math
import warnings
import logging

import numpy as np
import pandas as pd
import requests
from scipy.stats import norm

from config import (
    EXERCISE_TAX_RATE,
    get_commission_rate,
    get_exercise_fee_rate,
    get_symbol_kind,
    get_symbol_market,
)
from data.cleaner import DataCleaner
from data.downloader import MarketDownloader
from data.market_queue import (
    get_buy_queue_symbols,
    filter_by_option_symbols,
)
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


warnings.filterwarnings('ignore')
logger = logging.getLogger(__name__)


# ============================================================
# 🎯 پیکربندی Bale (کاربر باید وارد کند)
# ============================================================

BALE_ENABLED = True
BALE_BOT_TOKEN = "2144104837:M-1qbeWfXUTducuJS0aV4zu70P8xTr8Jzjk"
BALE_CHAT_ID = "@Option_Mehdi"
BALE_TOP_N = 1
BALE_PARSE_MODE = "Markdown"
BALE_TIMEOUT = 10

# --- جلوگیری از ارسال تکراری ---
BALE_DEDUP_ENABLED = True
BALE_SENT_HISTORY_FILE = current_dir / "sent_positions.json"


# ============================================================
# ثابت‌های Sentinel
# ============================================================
RISK_FREE_BREAK_EVEN_SENTINEL = -999.0
RISK_FREE_RETURN_SENTINEL = 999999.0
RISK_FREE_SCORE_SENTINEL = 999999.0
NEAR_EXPIRY_SCORE_BASE = RISK_FREE_SCORE_SENTINEL - 111111.0
NEAR_EXPIRY_PROB_SURVIVAL = 0.95


# ============================================================
# پارامترهای پیش‌فرض
# ============================================================

DEFAULT_MIN_MONTHLY_RETURN = 5.0
DEFAULT_MIN_MARGIN_FLOOR = 10.0
DEFAULT_MARGIN_PERCENTILE = 70
DEFAULT_MIN_RR_RATIO = 0.1
DEFAULT_W_R = 0.4
DEFAULT_W_M = 0.6
DEFAULT_SIGMA_FALLBACK = 0.30
DEFAULT_MIN_M_RISK = 2.0
DEFAULT_DECISION_THRESHOLD = 3.0

DEFAULT_VOL_QUALITY_POWER = 0.3
DEFAULT_IMBALANCE_POWER = 0.5

NEAR_EXPIRY_THRESHOLD_DAYS = 1.0
NEAR_EXPIRY_MIN_RETURN = 1.5
NEAR_EXPIRY_MIN_MARGIN = 2.0
NEAR_EXPIRY_W_R = 0.6
NEAR_EXPIRY_W_M = 0.4
T_DISPLAY_EPSILON = 0.02

VOLATILITY_FILE_NAME = "Historical_Volatility.xlsx"


# ==========================================================================================
# بخش ۰: Bale Notifier (داخلی)
# ==========================================================================================

def send_message_to_bale(message_text, bot_token=None, chat_id=None,
                          parse_mode=None, timeout=None):
    """ارسال پیام متنی به پیام‌رسان بله."""
    if not BALE_ENABLED:
        return None

    token = bot_token or BALE_BOT_TOKEN
    c_id = chat_id or BALE_CHAT_ID
    mode = parse_mode if parse_mode is not None else BALE_PARSE_MODE
    t_out = timeout or BALE_TIMEOUT

    if not token or token == "YOUR_BOT_TOKEN":
        print("[BALE] Bot token is not configured.")
        return None

    if not c_id or c_id == "YOUR_CHAT_ID":
        print("[BALE] Chat ID is not configured.")
        return None

    url = f'https://tapi.bale.ai/bot{token}/sendMessage'
    payload = {"chat_id": c_id, "text": message_text}
    if mode:
        payload["parse_mode"] = mode

    headers = {"Content-Type": "application/json"}

    try:
        response = requests.post(
            url,
            data=json.dumps(payload, ensure_ascii=False).encode('utf-8'),
            headers=headers,
            timeout=t_out
        )
        if response.status_code == 200:
            print(f"[BALE] Message sent successfully.")
            return response.json()
        else:
            print(f"[BALE] Failed: HTTP {response.status_code} - {response.text[:200]}")
            return None
    except requests.exceptions.Timeout:
        print("[BALE] Timeout occurred.")
        return None
    except Exception as e:
        print(f"[BALE] Error: {e}")
        return None


def _format_number(value, decimals=0):
    """فرمت عدد با کاما."""
    if value is None or pd.isna(value):
        return "N/A"
    try:
        if decimals > 0:
            return f"{float(value):,.{decimals}f}"
        return f"{int(round(float(value))):,}"
    except (ValueError, TypeError):
        return str(value)


def build_covered_call_message(row):
    """ساخت پیام Covered Call برای ارسال به بله."""
    underlying = row.get('underlying', 'N/A')
    option_sym = row.get('option_symbol', 'N/A')
    strike = row.get('strike', 0)
    premium = row.get('premium', 0)
    stock_price = row.get('stock_price', 0)
    break_even = row.get('break_even_price', 0)
    break_even_scale = row.get('break_even_percent_scale', 0)
    capital = row.get('capital_at_risk', 0)
    monthly_return = row.get('monthly_return_%', 0)
    score = row.get('composite_score', 0)
    days = row.get('days_to_maturity', 0)
    margin30 = row.get('margin30', 0)
    max_drop = row.get('max_drop_%', 0)
    raw_margin = row.get('raw_margin_percent', 0)
    maturity_date = row.get('maturity_date', 'N/A')

    message = (
        "🤖 *فرصت جدید استراتژی Covered Call*" + "\n"
        + "------------------------------------" + "\n"
        + f"📌 *نماد پایه:* `{underlying}`" + "\n"
        + f"🔴 *ساق فروش:* `{option_sym}` (اعمال: {_format_number(strike)} | پریمیوم: {_format_number(premium)})" + "\n"
        + "------------------------------------" + "\n"
        + f"📊 *قیمت سهم پایه:* `{_format_number(stock_price)}` ریال" + "\n"
        + f"🎯 *قیمت سربه‌سر:* `{_format_number(break_even)}` ریال" + "\n"
        + f"🛡 *حاشیه امنیت واقعی:* `%{_format_number(raw_margin, 2)}`" + "\n"
        + f"📐 *حاشیه امنیت ماهانه:* `%{_format_number(margin30, 2)}`" + "\n"
        + f"📏 *مقیاس سربه‌سر:* `%{_format_number(break_even_scale, 2)}`" + "\n"
        + f"📉 *حداکثر افت مجاز:* `%{_format_number(max_drop, 2)}`" + "\n"
        + f"💰 *سرمایه درگیر:* `{_format_number(capital)}` ریال" + "\n"
        + f"📈 *سود ماهانه:* `%{_format_number(monthly_return, 2)}`" + "\n"
        + f"🏆 *امتیاز نهایی:* `{_format_number(score, 1)}`" + "\n"
        + f"⏱ *روز تا سررسید:* `{days}` روز" + "\n"
        + f"📅 *تاریخ اعمال:* `{maturity_date}`"
    )

    return message


# ==========================================================================================
# بخش ۰.۵: مدیریت تاریخچه ارسال به بله (جلوگیری از تکرار)
# ==========================================================================================

def _load_sent_history():
    """بارگذاری تاریخچه ارسال‌شده از فایل JSON."""
    if not BALE_DEDUP_ENABLED:
        return {}

    if not BALE_SENT_HISTORY_FILE.exists():
        return {}

    try:
        with open(BALE_SENT_HISTORY_FILE, 'r', encoding='utf-8') as f:
            data = json.load(f)
        if not isinstance(data, dict):
            return {}
        return data
    except Exception as e:
        print(f"[BALE-DEDUP] Error loading history: {e}")
        return {}


def _save_sent_history(history):
    """ذخیره تاریخچه در فایل JSON."""
    if not BALE_DEDUP_ENABLED:
        return

    try:
        with open(BALE_SENT_HISTORY_FILE, 'w', encoding='utf-8') as f:
            json.dump(history, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"[BALE-DEDUP] Error saving history: {e}")


def _build_fingerprint(underlying, long_sym, short_sym, days, strategy):
    """
    ساخت اثر انگشت یکتا برای یک موقعیت.

    موقعیت جدید = تغییر در حداقل یکی از این ۵ عنصر:
        - استراتژی
        - نماد پایه
        - ساق خرید (Long) — در کاوردکال: "STOCK"
        - ساق فروش (Short)
        - روز تا سررسید

    Args:
        underlying: نماد پایه
        long_sym:   نماد ساق خرید (یا "STOCK" برای کاوردکال)
        short_sym:  نماد ساق فروش
        days:       روز تا سررسید
        strategy:   نام استراتژی (bull_call_spread یا covered_call)

    Returns:
        str: اثر انگشت یکتا
    """
    s = str(strategy).strip()
    u = str(underlying).strip()
    l = str(long_sym).strip()
    sh = str(short_sym).strip()
    d = str(int(days))
    return f"{s}|{u}|{l}|{sh}|{d}"


def _is_already_sent(fingerprint, history):
    """آیا این موقعیت قبلاً ارسال شده است؟"""
    return fingerprint in history


def _mark_as_sent(fingerprint, history, row, strategy_type="covered_call"):
    """ثبت موقعیت به‌عنوان ارسال‌شده."""
    history[fingerprint] = {
        'sent_at': pd.Timestamp.now().isoformat(),
        'strategy': strategy_type,
        'underlying': str(row.get('underlying', '')),
        'long_option_symbol': 'STOCK',
        'short_option_symbol': str(row.get('option_symbol', '')),
        'days_to_maturity': int(row.get('days_to_maturity', 0)),
        'maturity_date': str(row.get('maturity_date', '')),
        'composite_score': float(row.get('composite_score', 0))
            if pd.notna(row.get('composite_score', 0)) else 0.0,
    }


def send_best_position_to_bale(result_df):
    """
    ارسال بهترین موقعیت Covered Call به بله.

    منطق جلوگیری از تکرار:
        - اثر انگشت = strategy | underlying | "STOCK" | option_symbol | days
        - اگر موقعیت قبلاً ارسال شده باشد، دوباره ارسال نمی‌شود.
        - اگر حتی یکی از این ۵ عنصر تغییر کند، موقعیت جدید است و ارسال می‌شود.
    """
    if not BALE_ENABLED:
        print("\n[SKIP] Bale notification is disabled.")
        return

    if BALE_BOT_TOKEN == "YOUR_BOT_TOKEN" or BALE_CHAT_ID == "YOUR_CHAT_ID":
        print("\n[SKIP] Bale bot_token or chat_id is not configured.")
        return

    if result_df is None or result_df.empty:
        print("\n[BALE] No results to send.")
        return

    enter_df = result_df[result_df['decision'] == 'ENTER']
    if enter_df.empty:
        print("\n[BALE] No ENTER positions.")
        return

    # ===== بارگذاری تاریخچه =====
    history = _load_sent_history()
    print(f"\n[BALE-DEDUP] History: {len(history)} records.")

    top_n = enter_df.head(BALE_TOP_N)
    sent_count = 0
    skipped_count = 0

    for _, row in top_n.iterrows():
        # ساخت اثر انگشت (شامل strategy)
        fingerprint = _build_fingerprint(
            underlying=row.get('underlying', ''),
            long_sym="STOCK",
            short_sym=row.get('option_symbol', ''),
            days=row.get('days_to_maturity', 0),
            strategy="covered_call",
        )

        if _is_already_sent(fingerprint, history):
            print(f"[BALE-DEDUP] SKIP duplicate: {fingerprint}")
            skipped_count += 1
            continue

        # ===== ارسال =====
        msg = build_covered_call_message(row)
        result = send_message_to_bale(msg)

        if result is not None:
            _mark_as_sent(
                fingerprint, history, row,
                strategy_type="covered_call"
            )
            sent_count += 1
            print(f"[BALE-DEDUP] SENT new: {fingerprint}")

    # ===== ذخیره تاریخچه =====
    _save_sent_history(history)
    print(f"[BALE-DEDUP] Sent: {sent_count}, Skipped: {skipped_count}")


# ==========================================================================================
# بخش ۱: مدیریت صف
# ==========================================================================================

def normalize_symbol(symbol):
    """نرمال‌سازی ی/ک عربی و فارسی."""
    if not isinstance(symbol, str):
        return symbol
    return symbol.replace('ي', 'ی').replace('ك', 'ک')


def extract_option_symbols(df_options):
    """استخراج لیست نمادهای پایه دارای قرارداد اختیار از دیتای آپشن."""
    if df_options.empty:
        return []

    symbols = df_options['UnderlyingTicker'].dropna().unique().tolist()
    symbols = [normalize_symbol(s) for s in symbols]
    symbols = sorted(list(set(symbols)))

    return symbols


def filter_buy_queue_with_symbols(buy_queue, option_symbols):
    """فیلتر نمادهای صف خرید بر اساس option_symbols."""
    if not buy_queue:
        return []

    filtered = filter_by_option_symbols(buy_queue, option_symbols)

    if not filtered:
        print("  No option symbols in buy-queue.")
        return []

    buy_queue_symbols = [s['symbol'] for s in filtered]

    print(f"\n  === Buy-Queue Detection ===")
    print(f"  Total buy-queue symbols: {len(buy_queue)}")
    print(f"  Option symbols in buy-queue: {len(buy_queue_symbols)}")

    return buy_queue_symbols


def remove_buy_queue_underlyings(df_options, buy_queue_symbols):
    """حذف نمادهای پایه‌ای که در صف خرید هستند."""
    if df_options.empty:
        return df_options

    if not buy_queue_symbols:
        return df_options

    buy_queue_normalized = {normalize_symbol(s) for s in buy_queue_symbols}

    before_count = len(df_options)
    df_filtered = df_options[
        ~df_options['UnderlyingTicker'].apply(normalize_symbol).isin(
            buy_queue_normalized
        )
    ].copy()

    removed = before_count - len(df_filtered)
    print(f"  Removed: {removed} options")
    print(f"  Remaining: {len(df_filtered)} options")

    return df_filtered


# ==========================================================================================
# بخش ۲: محاسبات Covered Call
# ==========================================================================================

def covered_call_analysis(
    stock_price,
    strike_price,
    premium_call,
    contract_size,
    opt_sell_commission,
    stock_buy_commission,
    exercise_fee_rate,
    exercise_tax_rate,
    days,
):
    """محاسبه پارامترهای استراتژی Covered Call."""
    option_fee = -round(premium_call * contract_size * opt_sell_commission, 0)
    stock_buy_fee = -round(
        stock_price * contract_size * stock_buy_commission, 0
    )
    entry_fees = option_fee + stock_buy_fee

    exercise_fee = -round(
        strike_price * contract_size * exercise_fee_rate, 0
    )
    exercise_tax = -round(
        strike_price * contract_size * exercise_tax_rate, 0
    )

    premium_received = premium_call * contract_size
    stock_cost = -stock_price * contract_size

    net_investment = stock_cost + premium_received + entry_fees

    strike_received = strike_price * contract_size
    net_received = strike_received + exercise_fee + exercise_tax

    net_profit = net_received + net_investment

    if net_profit <= 0 or net_investment >= 0:
        return {'status': 'DISCARD'}

    capital_at_risk = max(1.0, abs(net_investment))
    days_raw = float(days)
    is_near_expiry = days_raw <= NEAR_EXPIRY_THRESHOLD_DAYS

    profit_percent = round((net_profit / capital_at_risk) * 100, 2)

    downside_protection = (
        premium_received + entry_fees + exercise_fee + exercise_tax
    )
    break_even_price = round(
        stock_price - (downside_protection / contract_size), 0
    )

    if stock_price > 0:
        break_even_percent = round(
            ((break_even_price - stock_price) / stock_price) * 100, 2
        )
    else:
        break_even_percent = 0.0

    max_drop_percent = round(
        ((stock_price - break_even_price) / stock_price) * 100, 2
    )

    raw_return_percent = profit_percent
    raw_margin_percent = max_drop_percent

    days_safe_for_display = max(T_DISPLAY_EPSILON, days_raw)
    time_factor = math.sqrt(days_safe_for_display / 30.0)

    break_even_percent_scale = round(break_even_percent / time_factor, 2)
    monthly_return = round(profit_percent * (30 / days_safe_for_display), 2)

    risk_reward_ratio = (
        round(net_profit / capital_at_risk, 2) if capital_at_risk > 0 else 0
    )

    return {
        'status': 'VALID',
        'is_near_expiry': is_near_expiry,
        'capital_at_risk': capital_at_risk,
        'net_profit': net_profit,
        'profit_percent': profit_percent,
        'monthly_return': monthly_return,
        'break_even_price': break_even_price,
        'break_even_percent': break_even_percent,
        'break_even_percent_scale': break_even_percent_scale,
        'max_drop_percent': max_drop_percent,
        'risk_reward_ratio': risk_reward_ratio,
        'raw_return_percent': raw_return_percent,
        'raw_margin_percent': raw_margin_percent,
    }


# ==========================================================================================
# بخش ۳: بارگذاری داده
# ==========================================================================================

def load_and_filter_data():
    """بارگذاری داده‌های بازار و فیلتر اولیه."""
    print("\n[1/3] Fetching buy-queue from TSETMC...")
    buy_queue = get_buy_queue_symbols()
    if not buy_queue:
        print("  WARNING: No buy-queue data received.")
        buy_queue = []
    else:
        print(f"  Buy-queue: {len(buy_queue)} symbols")

    print("\n[2/3] Downloading option market data...")
    df_raw = MarketDownloader.from_tsetmc_direct()
    df_cleaned = DataCleaner.clean(df_raw)
    df_final = DataCleaner.add_derived_columns(df_cleaned)

    is_call_mask = df_final['Type'].apply(
        lambda x: x.name == 'CALL' if hasattr(x, 'name')
        else str(x).upper() == 'CALL'
    )

    filter_option = df_final[
        (df_final['DaysToMaturity'] >= 0.0) & is_call_mask
    ].copy()

    print(f"  CALL options: {len(filter_option)}")

    print("\n[3/3] Filtering buy-queue by option symbols...")
    option_symbols = extract_option_symbols(filter_option)
    print(f"  Unique underlyings: {len(option_symbols)}")

    buy_queue_symbols = filter_buy_queue_with_symbols(
        buy_queue, option_symbols
    )

    filter_option = remove_buy_queue_underlyings(
        filter_option, buy_queue_symbols
    )

    return filter_option


def load_volatility_profile():
    """بارگذاری پروفایل نوسان از Historical_Volatility.xlsx."""
    filepath = current_dir / VOLATILITY_FILE_NAME

    if not filepath.exists():
        print(f"WARNING: Volatility file not found: {filepath}")
        return pd.DataFrame()

    try:
        df_vol = pd.read_excel(filepath)
        required_cols = ['UnderlyingTicker', 'HV_60']
        missing = [c for c in required_cols if c not in df_vol.columns]
        if missing:
            print(f"WARNING: Missing columns: {missing}")
            return pd.DataFrame()

        df_vol = df_vol.rename(columns={'UnderlyingTicker': 'underlying'})

        keep_cols = ['underlying', 'HV_60']
        if 'VolatilityQualityScore' in df_vol.columns:
            keep_cols.append('VolatilityQualityScore')

        return df_vol[keep_cols].copy()

    except Exception as e:
        print(f"ERROR loading volatility file: {e}")
        return pd.DataFrame()


# ==========================================================================================
# بخش ۴: افزودن M_30 و M_risk
# ==========================================================================================

def add_margin30_column(df):
    """استخراج M_30 از break_even_percent_scale."""
    df = df.copy()
    risk_free_mask = (
        df['break_even_percent_scale'] == RISK_FREE_BREAK_EVEN_SENTINEL
    )
    df['margin30'] = -df['break_even_percent_scale'].astype(float)
    df.loc[risk_free_mask, 'margin30'] = np.inf
    df['is_risk_free'] = risk_free_mask
    return df


def add_margin_risk_column(df):
    """محاسبه M_risk و Z_score با استفاده از HV_60 واقعی."""
    df = df.copy()
    if 'margin30' not in df.columns:
        df = add_margin30_column(df)

    risk_free_mask = df['is_risk_free']
    raw_margin_fraction = df['raw_margin_percent'].astype(float) / 100.0

    if 'HV_60' in df.columns:
        sigma = df['HV_60'].fillna(DEFAULT_SIGMA_FALLBACK).clip(
            lower=0.05, upper=2.0
        )
    else:
        sigma = pd.Series(DEFAULT_SIGMA_FALLBACK, index=df.index)

    T = df['days_to_maturity'].astype(float).clip(lower=T_DISPLAY_EPSILON)
    df['M_risk'] = raw_margin_fraction / (sigma * np.sqrt(T / 365.0))
    df.loc[risk_free_mask, 'M_risk'] = np.inf
    df['Z_score'] = df['M_risk']

    return df


# ==========================================================================================
# بخش ۵: فیلتر سخت
# ==========================================================================================

def apply_hard_constraints(
    df,
    min_monthly_return,
    min_margin_floor,
    margin_percentile,
    min_m_risk,
    near_expiry_min_return=NEAR_EXPIRY_MIN_RETURN,
    near_expiry_min_margin=NEAR_EXPIRY_MIN_MARGIN,
):
    """فیلتر سخت‌گیرانه."""
    if df.empty:
        return df

    df = add_margin30_column(df)

    risk_free_mask = df['is_risk_free']
    near_expiry_mask = df['is_near_expiry'].fillna(False) & (~risk_free_mask)
    normal_mask = ~(risk_free_mask | near_expiry_mask)

    df_rf = df[risk_free_mask].copy()
    df_near = df[near_expiry_mask].copy()
    df_normal = df[normal_mask].copy()

    if not df_near.empty:
        df_near = df_near[
            (df_near['raw_return_percent'] >= near_expiry_min_return)
            & (df_near['raw_margin_percent'] >= near_expiry_min_margin)
        ].copy()

    if not df_normal.empty:
        df_normal = df_normal[
            df_normal['monthly_return_%'] >= min_monthly_return
        ].copy()

        if not df_normal.empty:
            dynamic_threshold = np.percentile(
                df_normal['margin30'], margin_percentile
            )
            effective_margin_min = max(min_margin_floor, dynamic_threshold)

            df_normal = df_normal[
                df_normal['margin30'] >= effective_margin_min
            ].copy()

        if not df_normal.empty:
            df_normal = add_margin_risk_column(df_normal)
            df_normal = df_normal[df_normal['M_risk'] >= min_m_risk].copy()

    return pd.concat([df_rf, df_near, df_normal])


# ==========================================================================================
# بخش ۶: امتیازدهی
# ==========================================================================================

def calculate_composite_score(
    df,
    w_r=DEFAULT_W_R,
    w_m=DEFAULT_W_M,
    vol_quality_power=DEFAULT_VOL_QUALITY_POWER,
    imbalance_power=DEFAULT_IMBALANCE_POWER,
):
    """امتیازدهی نهایی با ضرایب ترکیبی."""
    if df.empty:
        return df

    df = df.copy()
    if 'margin30' not in df.columns:
        df = add_margin30_column(df)

    risk_free_mask = df['is_risk_free']
    near_expiry_mask = df['is_near_expiry'].fillna(False) & (~risk_free_mask)
    normal_mask = ~(risk_free_mask | near_expiry_mask)

    df_rf = df[risk_free_mask].copy()
    df_near = df[near_expiry_mask].copy()
    df_normal = df[normal_mask].copy()

    if not df_rf.empty:
        df_rf['prob_survival'] = 1.0
        df_rf['M_risk'] = np.inf
        df_rf['score_cobb_douglas'] = RISK_FREE_SCORE_SENTINEL
        df_rf['score_crra'] = RISK_FREE_SCORE_SENTINEL
        df_rf['vol_quality_factor'] = 1.0
        df_rf['imbalance_factor'] = 1.0
        df_rf['composite_score'] = RISK_FREE_SCORE_SENTINEL

    if not df_near.empty:
        raw_R = df_near['raw_return_percent'].astype(float)
        raw_M = df_near['raw_margin_percent'].astype(float)

        base_near = (
            NEAR_EXPIRY_SCORE_BASE
            + NEAR_EXPIRY_W_R * raw_R
            + NEAR_EXPIRY_W_M * raw_M
        ).round(4)

        df_near['prob_survival'] = NEAR_EXPIRY_PROB_SURVIVAL
        df_near['M_risk'] = np.inf
        df_near['score_cobb_douglas'] = base_near
        df_near['score_crra'] = base_near
        df_near['vol_quality_factor'] = 1.0
        df_near['imbalance_factor'] = 1.0
        df_near['composite_score'] = base_near

    if not df_normal.empty:
        if 'M_risk' not in df_normal.columns:
            df_normal = add_margin_risk_column(df_normal)

        R30 = df_normal['monthly_return_%'].astype(float).clip(lower=0.0)
        M30 = df_normal['margin30'].astype(float).clip(lower=0.0)
        M_risk = df_normal['M_risk'].astype(float)

        prob_survival = norm.cdf(M_risk)
        df_normal['prob_survival'] = np.round(prob_survival, 4)

        score_cd = (R30 ** w_r) * (M30 ** w_m)
        df_normal['score_cobb_douglas'] = (score_cd * prob_survival).round(4)

        score_crra = np.sqrt(R30 * M30) * prob_survival
        df_normal['score_crra'] = score_crra.round(4)

        if 'VolatilityQualityScore' in df_normal.columns:
            vq = df_normal['VolatilityQualityScore'].fillna(10.0).clip(
                lower=0.0, upper=10.0
            )
            vol_factor = (vq / 10.0) ** vol_quality_power
        else:
            vol_factor = pd.Series(1.0, index=df_normal.index)
        df_normal['vol_quality_factor'] = np.round(vol_factor, 4)

        ratio = (M30 / R30.replace(0, np.nan)).replace(
            [np.inf, -np.inf], np.nan
        ).fillna(1.0)
        imbalance_factor = np.minimum(1.0, ratio ** imbalance_power)
        df_normal['imbalance_factor'] = np.round(imbalance_factor, 4)

        df_normal['composite_score'] = (
            score_cd * prob_survival
            * vol_factor
            * imbalance_factor
        ).round(4)

    return pd.concat([df_rf, df_near, df_normal])


# ==========================================================================================
# بخش ۷: تابع اصلی
# ==========================================================================================

def analyze_covered_call(
    output_file: str = "covered_call_results.xlsx",
    min_monthly_return=DEFAULT_MIN_MONTHLY_RETURN,
    min_margin_floor=DEFAULT_MIN_MARGIN_FLOOR,
    margin_percentile=DEFAULT_MARGIN_PERCENTILE,
    min_rr_ratio=DEFAULT_MIN_RR_RATIO,
    w_r=DEFAULT_W_R,
    w_m=DEFAULT_W_M,
    vol_quality_power=DEFAULT_VOL_QUALITY_POWER,
    imbalance_power=DEFAULT_IMBALANCE_POWER,
    min_m_risk=DEFAULT_MIN_M_RISK,
    decision_threshold=DEFAULT_DECISION_THRESHOLD,
    near_expiry_min_return=NEAR_EXPIRY_MIN_RETURN,
    near_expiry_min_margin=NEAR_EXPIRY_MIN_MARGIN,
) -> pd.DataFrame:
    """اجرای کامل تحلیل Covered Call."""
    print("=" * 65)
    print("Covered Call Strategy")
    print("=" * 65)

    # ===== بارگذاری پروفایل نوسان =====
    df_vol = load_volatility_profile()
    if df_vol.empty:
        print("\nWARNING: Volatility profile not loaded.")

    # ===== بارگذاری داده بازار =====
    filter_option = load_and_filter_data()

    if filter_option.empty:
        print("No market data found.")
        return pd.DataFrame()

    # ===== محاسبه Covered Call =====
    results_fee = []

    for underlying_symbol, group in filter_option.groupby('UnderlyingTicker'):
        market = get_symbol_market(underlying_symbol)
        kind = get_symbol_kind(underlying_symbol)

        opt_sell_commission = get_commission_rate(market, 'option', False)
        stock_buy_commission = get_commission_rate(market, kind, True)
        exercise_fee_rate = get_exercise_fee_rate(market, kind)
        exercise_tax_rate = EXERCISE_TAX_RATE

        for _, item in group.iterrows():
            ticker = item['Ticker']
            strike_price = item['StrikePrice']
            premium_call = item['BidPrice']
            stock_price = item['UnderlyingPrice']
            contract_size = item['ContractSize']
            days = item['DaysToMaturity']

            if premium_call <= 0 or stock_price <= 0:
                continue

            res = covered_call_analysis(
                stock_price=stock_price,
                strike_price=strike_price,
                premium_call=premium_call,
                contract_size=contract_size,
                opt_sell_commission=opt_sell_commission,
                stock_buy_commission=stock_buy_commission,
                exercise_fee_rate=exercise_fee_rate,
                exercise_tax_rate=exercise_tax_rate,
                days=days,
            )

            if res['status'] == 'DISCARD':
                continue

            if res['risk_reward_ratio'] < min_rr_ratio:
                continue

            results_fee.append({
                'underlying': underlying_symbol,
                'is_near_expiry': res['is_near_expiry'],
                'stock_price': round(stock_price, 0),
                'option_symbol': ticker,
                'strike': strike_price,
                'premium': round(premium_call, 0),
                'contract_size': contract_size,
                'capital_at_risk': res['capital_at_risk'],
                'net_profit': res['net_profit'],
                'profit_percent': res['profit_percent'],
                'monthly_return_%': res['monthly_return'],
                'raw_return_percent': res['raw_return_percent'],
                'raw_margin_percent': res['raw_margin_percent'],
                'break_even_price': res['break_even_price'],
                'break_even_percent': res['break_even_percent'],
                'break_even_percent_scale': res['break_even_percent_scale'],
                'max_drop_%': res['max_drop_percent'],
                'risk_reward_ratio': res['risk_reward_ratio'],
                'days_to_maturity': days,
                'volume': int(item.get('Volume', 0)),
                'maturity_date': str(item.get('MaturityDate', '')),
            })

    if not results_fee:
        return pd.DataFrame()

    result_df = pd.DataFrame(results_fee)

    # ===== Merge با نوسان =====
    if df_vol is not None and not df_vol.empty:
        result_df = pd.merge(result_df, df_vol, on='underlying', how='left')
    else:
        result_df['HV_60'] = np.nan
        result_df['VolatilityQualityScore'] = np.nan

    # ===== فیلترینگ سخت =====
    result_df_filtered = apply_hard_constraints(
        result_df,
        min_monthly_return=min_monthly_return,
        min_margin_floor=min_margin_floor,
        margin_percentile=margin_percentile,
        min_m_risk=min_m_risk,
        near_expiry_min_return=near_expiry_min_return,
        near_expiry_min_margin=near_expiry_min_margin,
    )

    if result_df_filtered.empty:
        print("No valid setups after hard constraints.")
        return result_df_filtered

    # ===== امتیازدهی =====
    result_df_filtered = calculate_composite_score(
        result_df_filtered,
        w_r=w_r,
        w_m=w_m,
        vol_quality_power=vol_quality_power,
        imbalance_power=imbalance_power,
    )

    # ===== رتبه‌بندی =====
    result_df_filtered = result_df_filtered.sort_values(
        by='composite_score', ascending=False
    ).reset_index(drop=True)

    # ===== تصمیم =====
    always_enter_mask = (
        (result_df_filtered['composite_score'] == RISK_FREE_SCORE_SENTINEL)
        | (result_df_filtered['composite_score'] >= NEAR_EXPIRY_SCORE_BASE)
    )
    result_df_filtered['decision'] = np.where(
        always_enter_mask,
        'ENTER',
        np.where(
            result_df_filtered['composite_score'] > decision_threshold,
            'ENTER',
            'SKIP',
        ),
    )

    result_df_filtered['regime'] = np.select(
        [
            result_df_filtered['composite_score'] == RISK_FREE_SCORE_SENTINEL,
            result_df_filtered['composite_score'] >= NEAR_EXPIRY_SCORE_BASE,
        ],
        ['RISK_FREE', 'NEAR_EXPIRY'],
        default='NORMAL',
    )

    # چیدمان ستون‌ها
    column_order = [
        'underlying', 'regime', 'decision',
        'composite_score', 'score_cobb_douglas', 'score_crra',
        'prob_survival', 'M_risk', 'Z_score',
        'HV_60', 'VolatilityQualityScore',
        'vol_quality_factor', 'imbalance_factor',
        'stock_price', 'option_symbol', 'strike', 'premium',
        'contract_size', 'capital_at_risk', 'net_profit',
        'profit_percent', 'monthly_return_%', 'margin30',
        'raw_return_percent', 'raw_margin_percent',
        'break_even_price', 'break_even_percent',
        'break_even_percent_scale', 'max_drop_%',
        'risk_reward_ratio', 'days_to_maturity',
        'volume', 'maturity_date',
    ]
    column_order = [c for c in column_order if c in result_df_filtered.columns]
    result_df_filtered = result_df_filtered[column_order]

    # جایگزینی Sentinelها
    result_df_filtered['break_even_percent'] = result_df_filtered[
        'break_even_percent'
    ].replace(RISK_FREE_BREAK_EVEN_SENTINEL, 'Risk Free')
    result_df_filtered['break_even_percent_scale'] = result_df_filtered[
        'break_even_percent_scale'
    ].replace(RISK_FREE_BREAK_EVEN_SENTINEL, 'Risk Free')
    result_df_filtered['monthly_return_%'] = result_df_filtered[
        'monthly_return_%'
    ].replace(RISK_FREE_RETURN_SENTINEL, 'Infinite')
    result_df_filtered['margin30'] = result_df_filtered['margin30'].apply(
        lambda v: 'Risk Free' if v == np.inf else round(v, 2)
    )
    result_df_filtered['M_risk'] = result_df_filtered['M_risk'].apply(
        lambda v: 'N/A' if (pd.isna(v) or v == np.inf) else round(v, 2)
    )
    result_df_filtered['composite_score'] = result_df_filtered[
        'composite_score'
    ].replace(RISK_FREE_SCORE_SENTINEL, 'Risk Free')

    # ===== ذخیره در اکسل =====
    save_results_to_excel(result_df_filtered, output_file)

    # ===== نمایش خلاصه =====
    print("\n=== Decision Summary ===")
    for regime in ['RISK_FREE', 'NEAR_EXPIRY', 'NORMAL']:
        regime_df = result_df_filtered[result_df_filtered['regime'] == regime]
        if regime_df.empty:
            continue
        enter_count = (regime_df['decision'] == 'ENTER').sum()
        skip_count = (regime_df['decision'] == 'SKIP').sum()
        print(f"   [{regime}] ENTER: {enter_count}   SKIP: {skip_count}")

    entered = result_df_filtered[result_df_filtered['decision'] == 'ENTER']
    if not entered.empty:
        best = entered.iloc[0]
        print("\n=== Best Position ===")
        print(f"   Composite Score: {best['composite_score']}")

    # ===== ارسال به بله =====
    send_best_position_to_bale(result_df_filtered)

    return result_df_filtered


# ==========================================================================================
# بخش ۸: ذخیره در اکسل
# ==========================================================================================

def save_results_to_excel(result_df, filename="covered_call_results.xlsx"):
    """ذخیره نتایج خروجی در فایل اکسل."""
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(
        start_color='203764', end_color='203764', fill_type='solid'
    )
    alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    body_font = Font(name='Segoe UI', size=10)
    gray_font = Font(color='808080', italic=True, name='Segoe UI', size=10)

    rename_dict = {
        'underlying': 'نماد پایه',
        'regime': 'رژیم',
        'decision': 'تصمیم',
        'composite_score': 'امتیاز نهایی',
        'score_cobb_douglas': 'امتیاز Cobb-Douglas',
        'score_crra': 'امتیاز CRRA',
        'prob_survival': 'احتمال بقا',
        'M_risk': 'M_risk',
        'Z_score': 'Z-Score',
        'HV_60': 'نوسان ۶۰ روزه',
        'VolatilityQualityScore': 'کیفیت نوسان',
        'vol_quality_factor': 'ضریب کیفیت نوسان',
        'imbalance_factor': 'ضریب عدم‌تعادل',
        'stock_price': 'قیمت سهم',
        'option_symbol': 'نماد اختیار',
        'strike': 'قیمت اعمال',
        'premium': 'پریمیوم فروش',
        'contract_size': 'اندازه قرارداد',
        'capital_at_risk': 'سرمایه درگیر',
        'net_profit': 'سود خالص',
        'profit_percent': 'درصد سود',
        'monthly_return_%': 'سود ماهانه (R30)',
        'margin30': 'حاشیه ماهانه (M30)',
        'raw_return_percent': 'سود خام',
        'raw_margin_percent': 'حاشیه خام',
        'break_even_price': 'قیمت سربه‌سر',
        'break_even_percent': 'درصد رشد تا سربه‌سر',
        'break_even_percent_scale': 'مقیاس سربه‌سر',
        'max_drop_%': 'حداکثر افت مجاز',
        'risk_reward_ratio': 'R/R',
        'days_to_maturity': 'روز تا سررسید',
        'volume': 'حجم',
        'maturity_date': 'تاریخ اعمال',
    }

    result_df_renamed = result_df.rename(columns=rename_dict)
    filepath = current_dir / filename

    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        result_df_renamed.to_excel(
            writer, sheet_name='covered_call', index=False
        )
        worksheet = writer.sheets['covered_call']

        for col_idx in range(1, len(result_df_renamed.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = alignment

        for row_idx in range(2, len(result_df_renamed) + 2):
            for col_idx in range(1, len(result_df_renamed.columns) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                val = cell.value
                if val is None or pd.isna(val) or val == "":
                    cell.value = "-"
                    cell.font = gray_font
                else:
                    cell.font = body_font
                cell.alignment = alignment

        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(result_df_renamed.columns))}"
            f"{len(result_df_renamed) + 1}"
        )
        worksheet.freeze_panes = 'A2'

        for col in worksheet.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                if cell.value:
                    text = str(cell.value)
                    if '\n' in text:
                        lines = text.split('\n')
                        line_length = max(len(line) for line in lines)
                    else:
                        line_length = len(text)
                    if line_length > max_length:
                        max_length = line_length
            adjusted_width = min(max_length + 5, 50)
            worksheet.column_dimensions[column].width = adjusted_width

    print(f"Results saved to: {filename}")


# ==========================================================================================
# بخش ۹: اجرای مستقیم
# ==========================================================================================

if __name__ == "__main__":
    OUTPUT_FILE = "covered_call_results.xlsx"

    df_result = analyze_covered_call(output_file=OUTPUT_FILE,)