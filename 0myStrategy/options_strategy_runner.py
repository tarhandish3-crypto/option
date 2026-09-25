# 0myStrategy/strategies_analyzer.py
# -*- coding: utf-8 -*-
"""
ماژول ترکیبی استراتژی‌های Bull Call Spread و Covered Call.

معماری این ماژول:

    1) دریافت صف نمادهای پایه از TSETMC (یک‌بار)
    2) دانلود دیتای آپشن (یک‌بار، تازه‌ترین)
    3) فیلتر CALL
    4) استخراج option_symbols از UnderlyingTicker
    5) حذف نمادهای صف خرید
    6) بارگذاری Historical_Volatility.xlsx
    7) اجرای Bull Call Spread (اختیاری - با پارامترهای اختصاصی)
    8) اجرای Covered Call (اختیاری - با پارامترهای اختصاصی)
    9) ذخیره در یک فایل اکسل با دو شیت مجزا

نکات کلیدی:
    - دیتای آپشن فقط یک‌بار دانلود می‌شود (اشتراک بین دو استراتژی).
    - صف نمادهای پایه فقط یک‌بار دریافت می‌شود.
    - پارامترهای هر استراتژی مستقل هستند.
    - فعال/غیرفعال بودن هر استراتژی به اختیار کاربر.
    - خروجی: strategy_results.xlsx با دو شیت.
    - ستون‌های اضافی (Cobb-Douglas, CRRA, prob_survival, Z_score,
      vol_quality_factor, imbalance_factor) در اکسل ذخیره نمی‌شوند.
    - ستون‌های محاسباتی (raw_return_percent, raw_margin_percent,
      break_even_percent_scale) حفظ می‌شوند.
"""

import sys
from pathlib import Path

current_file_path = Path(__file__).resolve()
current_dir = current_file_path.parent
root_dir = current_dir.parent
sys.path.append(str(root_dir))

import math
import warnings
import logging

import numpy as np
import pandas as pd
from scipy.stats import norm

from config import (
    EXERCISE_TAX_RATE,
    get_commission_rate,
    get_exercise_fee_rate,
    get_symbol_kind,
    get_symbol_market,
)
from data.cleaner import DataCleaner
from data.downloader import MarketDownloader
from data.market_queue import (
    get_buy_queue_symbols,
    filter_by_option_symbols,
)
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


warnings.filterwarnings('ignore')
logger = logging.getLogger(__name__)


# ============================================================
# ثابت‌های Sentinel
# ============================================================
RISK_FREE_BREAK_EVEN_SENTINEL = -999.0
RISK_FREE_RETURN_SENTINEL = 999999.0
RISK_FREE_SCORE_SENTINEL = 999999.0
NEAR_EXPIRY_SCORE_BASE = RISK_FREE_SCORE_SENTINEL - 111111.0
NEAR_EXPIRY_PROB_SURVIVAL = 0.95


# ============================================================
# پارامترهای اختصاصی Bull Call Spread
# ============================================================

BCS_ENABLED = True                      # فعال/غیرفعال
BCS_MIN_MONTHLY_RETURN = 7.0           # حداقل سود ماهانه (%)
BCS_MIN_MARGIN_FLOOR = 20.0            # کف مطلق حاشیه (%)
BCS_MARGIN_PERCENTILE = 60             # صدک آستانه پویا
BCS_MIN_RR_RATIO = 0.1                 # حداقل R/R
BCS_W_R = 0.4                          # وزن سود
BCS_W_M = 0.6                          # وزن امنیت
BCS_MIN_M_RISK = 2.0                   # حداقل Z-Score
BCS_DECISION_THRESHOLD = 3.0           # آستانه تصمیم
BCS_VOL_QUALITY_POWER = 0.3
BCS_IMBALANCE_POWER = 0.5


# ============================================================
# پارامترهای اختصاصی Covered Call
# ============================================================

CC_ENABLED = True                       # فعال/غیرفعال
CC_MIN_MONTHLY_RETURN = 5.0            # حداقل سود ماهانه (%)
CC_MIN_MARGIN_FLOOR = 10.0             # کف مطلق حاشیه (%)
CC_MARGIN_PERCENTILE = 70              # صدک آستانه پویا
CC_MIN_RR_RATIO = 0.1                  # حداقل R/R
CC_W_R = 0.4                           # وزن سود
CC_W_M = 0.6                           # وزن امنیت
CC_MIN_M_RISK = 2.0                    # حداقل Z-Score
CC_DECISION_THRESHOLD = 3.0            # آستانه تصمیم
CC_VOL_QUALITY_POWER = 0.3
CC_IMBALANCE_POWER = 0.5


# ============================================================
# پارامترهای مشترک
# ============================================================

DEFAULT_SIGMA_FALLBACK = 0.30
NEAR_EXPIRY_THRESHOLD_DAYS = 1.0
NEAR_EXPIRY_MIN_RETURN = 1.5
NEAR_EXPIRY_MIN_MARGIN = 2.0
NEAR_EXPIRY_W_R = 0.6
NEAR_EXPIRY_W_M = 0.4
T_DISPLAY_EPSILON = 0.02

VOLATILITY_FILE_NAME = "Historical_Volatility.xlsx"
OUTPUT_FILE_NAME = "strategy_results.xlsx"


# ==========================================================================================
# بخش ۱: مدیریت صف
# ==========================================================================================

def normalize_symbol(symbol):
    """نرمال‌سازی ی/ک عربی و فارسی."""
    if not isinstance(symbol, str):
        return symbol
    return symbol.replace('ي', 'ی').replace('ك', 'ک')


def extract_option_symbols(df_options):
    """استخراج لیست نمادهای پایه دارای قرارداد اختیار از دیتای آپشن."""
    if df_options.empty:
        return []

    symbols = df_options['UnderlyingTicker'].dropna().unique().tolist()
    symbols = [normalize_symbol(s) for s in symbols]
    symbols = sorted(list(set(symbols)))

    return symbols


def filter_buy_queue_with_symbols(buy_queue, option_symbols):
    """فیلتر نمادهای صف خرید بر اساس option_symbols."""
    if not buy_queue:
        return []

    filtered = filter_by_option_symbols(buy_queue, option_symbols)

    if not filtered:
        print("  No option symbols in buy-queue.")
        return []

    buy_queue_symbols = [s['symbol'] for s in filtered]

    print(f"\n  === Buy-Queue Detection ===")
    print(f"  Total buy-queue symbols: {len(buy_queue)}")
    print(f"  Option symbols in buy-queue: {len(buy_queue_symbols)}")

    return buy_queue_symbols


def remove_buy_queue_underlyings(df_options, buy_queue_symbols):
    """حذف نمادهای پایه‌ای که در صف خرید هستند."""
    if df_options.empty:
        return df_options

    if not buy_queue_symbols:
        return df_options

    buy_queue_normalized = {normalize_symbol(s) for s in buy_queue_symbols}

    before_count = len(df_options)
    df_filtered = df_options[
        ~df_options['UnderlyingTicker'].apply(normalize_symbol).isin(
            buy_queue_normalized
        )
    ].copy()

    removed = before_count - len(df_filtered)
    print(f"  Removed: {removed} options")
    print(f"  Remaining: {len(df_filtered)} options")

    return df_filtered


# ==========================================================================================
# بخش ۲: بارگذاری داده (مشترک)
# ==========================================================================================

def load_and_filter_data():
    """بارگذاری داده‌های بازار و فیلتر اولیه (مشترک)."""
    print("\n[1/3] Fetching buy-queue from TSETMC...")
    buy_queue = get_buy_queue_symbols()
    if not buy_queue:
        print("  WARNING: No buy-queue data received.")
        buy_queue = []
    else:
        print(f"  Buy-queue: {len(buy_queue)} symbols")

    print("\n[2/3] Downloading option market data...")
    df_raw = MarketDownloader.from_tsetmc_direct()
    df_cleaned = DataCleaner.clean(df_raw)
    df_final = DataCleaner.add_derived_columns(df_cleaned)

    is_call_mask = df_final['Type'].apply(
        lambda x: x.name == 'CALL' if hasattr(x, 'name')
        else str(x).upper() == 'CALL'
    )

    filter_option = df_final[
        (df_final['DaysToMaturity'] >= 0.0) & is_call_mask
    ].copy()

    print(f"  CALL options: {len(filter_option)}")

    print("\n[3/3] Filtering buy-queue by option symbols...")
    option_symbols = extract_option_symbols(filter_option)
    print(f"  Unique underlyings: {len(option_symbols)}")

    buy_queue_symbols = filter_buy_queue_with_symbols(
        buy_queue, option_symbols
    )

    filter_option = remove_buy_queue_underlyings(
        filter_option, buy_queue_symbols
    )

    return filter_option


def load_volatility_profile():
    """بارگذاری پروفایل نوسان از Historical_Volatility.xlsx."""
    filepath = current_dir / VOLATILITY_FILE_NAME

    if not filepath.exists():
        print(f"WARNING: Volatility file not found: {filepath}")
        return pd.DataFrame()

    try:
        df_vol = pd.read_excel(filepath)
        required_cols = ['UnderlyingTicker', 'HV_60']
        missing = [c for c in required_cols if c not in df_vol.columns]
        if missing:
            print(f"WARNING: Missing columns: {missing}")
            return pd.DataFrame()

        df_vol = df_vol.rename(columns={'UnderlyingTicker': 'underlying'})

        keep_cols = ['underlying', 'HV_60']
        if 'VolatilityQualityScore' in df_vol.columns:
            keep_cols.append('VolatilityQualityScore')

        return df_vol[keep_cols].copy()

    except Exception as e:
        print(f"ERROR loading volatility file: {e}")
        return pd.DataFrame()


# ==========================================================================================
# بخش ۳: توابع مشترک M_30 و M_risk
# ==========================================================================================

def add_margin30_column(df):
    """استخراج M_30 از break_even_percent_scale."""
    df = df.copy()
    risk_free_mask = (
        df['break_even_percent_scale'] == RISK_FREE_BREAK_EVEN_SENTINEL
    )
    df['margin30'] = -df['break_even_percent_scale'].astype(float)
    df.loc[risk_free_mask, 'margin30'] = np.inf
    df['is_risk_free'] = risk_free_mask
    return df


def add_margin_risk_column(df):
    """محاسبه M_risk و Z_score با استفاده از HV_60 واقعی."""
    df = df.copy()
    if 'margin30' not in df.columns:
        df = add_margin30_column(df)

    risk_free_mask = df['is_risk_free']
    raw_margin_fraction = df['raw_margin_percent'].astype(float) / 100.0

    if 'HV_60' in df.columns:
        sigma = df['HV_60'].fillna(DEFAULT_SIGMA_FALLBACK).clip(
            lower=0.05, upper=2.0
        )
    else:
        sigma = pd.Series(DEFAULT_SIGMA_FALLBACK, index=df.index)

    T = df['days_to_maturity'].astype(float).clip(lower=T_DISPLAY_EPSILON)
    df['M_risk'] = raw_margin_fraction / (sigma * np.sqrt(T / 365.0))
    df.loc[risk_free_mask, 'M_risk'] = np.inf
    df['Z_score'] = df['M_risk']

    return df


# ==========================================================================================
# بخش ۴: محاسبات Bull Call Spread
# ==========================================================================================

def bull_call_spread_analysis(
    stock_price,
    long_strike,
    long_ask_premium,
    short_strike,
    short_bid_premium,
    contract_size,
    opt_buy_commission,
    opt_sell_commission,
    exercise_fee_rate,
    exercise_tax_rate,
    days,
):
    """محاسبه پارامترهای Bull Call Spread."""
    long_premium_total = round(long_ask_premium * contract_size, 0)
    long_entry_fee = round(long_premium_total * opt_buy_commission, 0)

    short_premium_total = round(short_bid_premium * contract_size, 0)
    short_entry_fee = -round(short_premium_total * opt_sell_commission, 0)

    net_debit = (long_premium_total + long_entry_fee) - (
        short_premium_total + short_entry_fee
    )

    long_exercise_fee = round((long_strike * contract_size) * exercise_fee_rate, 0)
    short_exercise_fee = round((short_strike * contract_size) * exercise_fee_rate, 0)
    short_transfer_tax = round((short_strike * contract_size) * exercise_tax_rate, 0)

    short_total_exercise_cost = short_exercise_fee + short_transfer_tax
    total_exercise_costs = long_exercise_fee + short_total_exercise_cost

    min_profit_or_loss = -net_debit
    max_payoff = (short_strike - long_strike) * contract_size
    max_net_profit = max_payoff - net_debit - total_exercise_costs

    if max_net_profit <= 0 or net_debit <= 0:
        return {'status': 'DISCARD'}

    capital_at_risk = max(1.0, net_debit)
    days_raw = float(days)
    is_near_expiry = days_raw <= NEAR_EXPIRY_THRESHOLD_DAYS

    if min_profit_or_loss >= 0:
        return {
            'status': 'RISK_FREE',
            'is_near_expiry': False,
            'capital_at_risk': 0,
            'max_net_profit': max_net_profit,
            'max_profit_percent': 'Arbitrage',
            'monthly_return': RISK_FREE_RETURN_SENTINEL,
            'break_even_price': 'Risk Free',
            'break_even_percent': RISK_FREE_BREAK_EVEN_SENTINEL,
            'break_even_percent_scale': RISK_FREE_BREAK_EVEN_SENTINEL,
            'risk_reward_ratio': 'Infinite',
            'raw_return_percent': None,
            'raw_margin_percent': None,
        }

    break_even_price = long_strike + ((net_debit + long_exercise_fee) / contract_size)

    if stock_price > 0:
        break_even_percent = round(
            ((break_even_price - stock_price) / stock_price) * 100, 2
        )
    else:
        break_even_percent = 0.0

    max_profit_percent = round((max_net_profit / capital_at_risk) * 100, 2)
    raw_return_percent = max_profit_percent
    raw_margin_percent = -break_even_percent
    risk_reward_ratio = round(max_net_profit / capital_at_risk, 2)

    days_safe_for_display = max(T_DISPLAY_EPSILON, days_raw)
    time_factor = math.sqrt(days_safe_for_display / 30.0)

    break_even_percent_scale = round(break_even_percent / time_factor, 2)
    monthly_return = round(max_profit_percent * (30 / days_safe_for_display), 2)

    return {
        'status': 'VALID',
        'is_near_expiry': is_near_expiry,
        'capital_at_risk': capital_at_risk,
        'max_net_profit': max_net_profit,
        'max_profit_percent': max_profit_percent,
        'monthly_return': monthly_return,
        'break_even_price': round(break_even_price, 0),
        'break_even_percent': break_even_percent,
        'break_even_percent_scale': break_even_percent_scale,
        'risk_reward_ratio': risk_reward_ratio,
        'raw_return_percent': raw_return_percent,
        'raw_margin_percent': raw_margin_percent,
    }


# ==========================================================================================
# بخش ۵: محاسبات Covered Call
# ==========================================================================================

def covered_call_analysis(
    stock_price,
    strike_price,
    premium_call,
    contract_size,
    opt_sell_commission,
    stock_buy_commission,
    exercise_fee_rate,
    exercise_tax_rate,
    days,
):
    """محاسبه پارامترهای Covered Call."""
    option_fee = -round(premium_call * contract_size * opt_sell_commission, 0)
    stock_buy_fee = -round(
        stock_price * contract_size * stock_buy_commission, 0
    )
    entry_fees = option_fee + stock_buy_fee

    exercise_fee = -round(
        strike_price * contract_size * exercise_fee_rate, 0
    )
    exercise_tax = -round(
        strike_price * contract_size * exercise_tax_rate, 0
    )

    premium_received = premium_call * contract_size
    stock_cost = -stock_price * contract_size

    net_investment = stock_cost + premium_received + entry_fees

    strike_received = strike_price * contract_size
    net_received = strike_received + exercise_fee + exercise_tax

    net_profit = net_received + net_investment

    if net_profit <= 0 or net_investment >= 0:
        return {'status': 'DISCARD'}

    capital_at_risk = max(1.0, abs(net_investment))
    days_raw = float(days)
    is_near_expiry = days_raw <= NEAR_EXPIRY_THRESHOLD_DAYS

    profit_percent = round((net_profit / capital_at_risk) * 100, 2)

    downside_protection = (
        premium_received + entry_fees + exercise_fee + exercise_tax
    )
    break_even_price = round(
        stock_price - (downside_protection / contract_size), 0
    )

    if stock_price > 0:
        break_even_percent = round(
            ((break_even_price - stock_price) / stock_price) * 100, 2
        )
    else:
        break_even_percent = 0.0

    max_drop_percent = round(
        ((stock_price - break_even_price) / stock_price) * 100, 2
    )

    raw_return_percent = profit_percent
    raw_margin_percent = max_drop_percent

    days_safe_for_display = max(T_DISPLAY_EPSILON, days_raw)
    time_factor = math.sqrt(days_safe_for_display / 30.0)

    break_even_percent_scale = round(break_even_percent / time_factor, 2)
    monthly_return = round(profit_percent * (30 / days_safe_for_display), 2)

    risk_reward_ratio = (
        round(net_profit / capital_at_risk, 2) if capital_at_risk > 0 else 0
    )

    return {
        'status': 'VALID',
        'is_near_expiry': is_near_expiry,
        'capital_at_risk': capital_at_risk,
        'net_profit': net_profit,
        'profit_percent': profit_percent,
        'monthly_return': monthly_return,
        'break_even_price': break_even_price,
        'break_even_percent': break_even_percent,
        'break_even_percent_scale': break_even_percent_scale,
        'max_drop_percent': max_drop_percent,
        'risk_reward_ratio': risk_reward_ratio,
        'raw_return_percent': raw_return_percent,
        'raw_margin_percent': raw_margin_percent,
    }


# ==========================================================================================
# بخش ۶: فیلتر سخت و امتیازدهی (مشترک، پارامتری)
# ==========================================================================================

def apply_hard_constraints(
    df,
    min_monthly_return,
    min_margin_floor,
    margin_percentile,
    min_m_risk,
    near_expiry_min_return=NEAR_EXPIRY_MIN_RETURN,
    near_expiry_min_margin=NEAR_EXPIRY_MIN_MARGIN,
):
    """فیلتر سخت‌گیرانه (پارامتری برای هر استراتژی)."""
    if df.empty:
        return df

    df = add_margin30_column(df)

    risk_free_mask = df['is_risk_free']
    near_expiry_mask = df['is_near_expiry'].fillna(False) & (~risk_free_mask)
    normal_mask = ~(risk_free_mask | near_expiry_mask)

    df_rf = df[risk_free_mask].copy()
    df_near = df[near_expiry_mask].copy()
    df_normal = df[normal_mask].copy()

    if not df_near.empty:
        df_near = df_near[
            (df_near['raw_return_percent'] >= near_expiry_min_return)
            & (df_near['raw_margin_percent'] >= near_expiry_min_margin)
        ].copy()

    if not df_normal.empty:
        df_normal = df_normal[
            df_normal['monthly_return_%'] >= min_monthly_return
        ].copy()

        if not df_normal.empty:
            dynamic_threshold = np.percentile(
                df_normal['margin30'], margin_percentile
            )
            effective_margin_min = max(min_margin_floor, dynamic_threshold)

            df_normal = df_normal[
                df_normal['margin30'] >= effective_margin_min
            ].copy()

        if not df_normal.empty:
            df_normal = add_margin_risk_column(df_normal)
            df_normal = df_normal[df_normal['M_risk'] >= min_m_risk].copy()

    return pd.concat([df_rf, df_near, df_normal])


def calculate_composite_score(
    df,
    w_r,
    w_m,
    vol_quality_power,
    imbalance_power,
):
    """امتیازدهی نهایی (پارامتری برای هر استراتژی)."""
    if df.empty:
        return df

    df = df.copy()
    if 'margin30' not in df.columns:
        df = add_margin30_column(df)

    risk_free_mask = df['is_risk_free']
    near_expiry_mask = df['is_near_expiry'].fillna(False) & (~risk_free_mask)
    normal_mask = ~(risk_free_mask | near_expiry_mask)

    df_rf = df[risk_free_mask].copy()
    df_near = df[near_expiry_mask].copy()
    df_normal = df[normal_mask].copy()

    if not df_rf.empty:
        df_rf['prob_survival'] = 1.0
        df_rf['M_risk'] = np.inf
        df_rf['composite_score'] = RISK_FREE_SCORE_SENTINEL

    if not df_near.empty:
        raw_R = df_near['raw_return_percent'].astype(float)
        raw_M = df_near['raw_margin_percent'].astype(float)

        base_near = (
            NEAR_EXPIRY_SCORE_BASE
            + NEAR_EXPIRY_W_R * raw_R
            + NEAR_EXPIRY_W_M * raw_M
        ).round(4)

        df_near['prob_survival'] = NEAR_EXPIRY_PROB_SURVIVAL
        df_near['M_risk'] = np.inf
        df_near['composite_score'] = base_near

    if not df_normal.empty:
        if 'M_risk' not in df_normal.columns:
            df_normal = add_margin_risk_column(df_normal)

        R30 = df_normal['monthly_return_%'].astype(float).clip(lower=0.0)
        M30 = df_normal['margin30'].astype(float).clip(lower=0.0)
        M_risk = df_normal['M_risk'].astype(float)

        prob_survival = norm.cdf(M_risk)
        df_normal['prob_survival'] = np.round(prob_survival, 4)

        score_cd = (R30 ** w_r) * (M30 ** w_m)

        if 'VolatilityQualityScore' in df_normal.columns:
            vq = df_normal['VolatilityQualityScore'].fillna(10.0).clip(
                lower=0.0, upper=10.0
            )
            vol_factor = (vq / 10.0) ** vol_quality_power
        else:
            vol_factor = pd.Series(1.0, index=df_normal.index)

        ratio = (M30 / R30.replace(0, np.nan)).replace(
            [np.inf, -np.inf], np.nan
        ).fillna(1.0)
        imbalance_factor = np.minimum(1.0, ratio ** imbalance_power)

        df_normal['composite_score'] = (
            score_cd * prob_survival
            * vol_factor
            * imbalance_factor
        ).round(4)

    return pd.concat([df_rf, df_near, df_normal])


# ==========================================================================================
# بخش ۷: اجرای Bull Call Spread
# ==========================================================================================

def run_bull_call_spread(df_options, df_vol):
    """اجرای کامل استراتژی Bull Call Spread."""
    print("\n" + "=" * 65)
    print("Running Bull Call Spread Strategy")
    print("=" * 65)

    results_fee = []

    for underlying_symbol, group in df_options.groupby('UnderlyingTicker'):
        market = get_symbol_market(underlying_symbol)
        kind = get_symbol_kind(underlying_symbol)

        opt_buy_commission = get_commission_rate(market, 'option', True)
        opt_sell_commission = get_commission_rate(market, 'option', False)
        exercise_fee_rate = get_exercise_fee_rate(market, kind)

        for days, sub_group in group.groupby('DaysToMaturity'):
            sub_group_sorted = sub_group.sort_values('StrikePrice')
            options_list = sub_group_sorted.to_dict('records')
            n = len(options_list)

            for i in range(n):
                for j in range(i + 1, n):
                    long_leg = options_list[i]
                    short_leg = options_list[j]

                    stock_price = long_leg['UnderlyingPrice']
                    contract_size = long_leg['ContractSize']

                    long_ask = long_leg.get('AskPrice', 0)
                    short_bid = short_leg.get('BidPrice', 0)

                    if (
                        pd.isna(long_ask)
                        or long_ask <= 0
                        or pd.isna(short_bid)
                        or short_bid <= 0
                    ):
                        continue

                    res = bull_call_spread_analysis(
                        stock_price=stock_price,
                        long_strike=long_leg['StrikePrice'],
                        long_ask_premium=long_ask,
                        short_strike=short_leg['StrikePrice'],
                        short_bid_premium=short_bid,
                        contract_size=contract_size,
                        opt_buy_commission=opt_buy_commission,
                        opt_sell_commission=opt_sell_commission,
                        exercise_fee_rate=exercise_fee_rate,
                        exercise_tax_rate=EXERCISE_TAX_RATE,
                        days=days,
                    )

                    if res['status'] == 'DISCARD':
                        continue

                    if (
                        res['status'] != 'RISK_FREE'
                        and res['risk_reward_ratio'] < BCS_MIN_RR_RATIO
                    ):
                        continue

                    results_fee.append({
                        'underlying': underlying_symbol,
                        'is_near_expiry': res['is_near_expiry'],
                        'stock_price': round(stock_price, 0),
                        'long_option_symbol': long_leg['Ticker'],
                        'long_strike': long_leg['StrikePrice'],
                        'long_ask_premium': round(long_ask, 0),
                        'short_option_symbol': short_leg['Ticker'],
                        'short_strike': short_leg['StrikePrice'],
                        'short_bid_premium': round(short_bid, 0),
                        'capital_at_risk': res['capital_at_risk'],
                        'max_net_profit': res['max_net_profit'],
                        'max_profit_percent': res['max_profit_percent'],
                        'monthly_return_%': res['monthly_return'],
                        'raw_return_percent': res['raw_return_percent'],
                        'raw_margin_percent': res['raw_margin_percent'],
                        'break_even_price': res['break_even_price'],
                        'break_even_percent': res['break_even_percent'],
                        'break_even_percent_scale': res['break_even_percent_scale'],
                        'risk_reward_ratio': res['risk_reward_ratio'],
                        'days_to_maturity': days,
                        'long_volume': int(long_leg.get('Volume', 0)),
                        'short_volume': int(short_leg.get('Volume', 0)),
                    })

    if not results_fee:
        print("No valid Bull Call Spread setups.")
        return pd.DataFrame()

    result_df = pd.DataFrame(results_fee)

    # Merge با نوسان
    if df_vol is not None and not df_vol.empty:
        result_df = pd.merge(result_df, df_vol, on='underlying', how='left')
    else:
        result_df['HV_60'] = np.nan
        result_df['VolatilityQualityScore'] = np.nan

    # فیلترینگ سخت
    result_df_filtered = apply_hard_constraints(
        result_df,
        min_monthly_return=BCS_MIN_MONTHLY_RETURN,
        min_margin_floor=BCS_MIN_MARGIN_FLOOR,
        margin_percentile=BCS_MARGIN_PERCENTILE,
        min_m_risk=BCS_MIN_M_RISK,
    )

    if result_df_filtered.empty:
        print("No valid setups after hard constraints.")
        return result_df_filtered

    # امتیازدهی
    result_df_filtered = calculate_composite_score(
        result_df_filtered,
        w_r=BCS_W_R,
        w_m=BCS_W_M,
        vol_quality_power=BCS_VOL_QUALITY_POWER,
        imbalance_power=BCS_IMBALANCE_POWER,
    )

    # رتبه‌بندی
    result_df_filtered = result_df_filtered.sort_values(
        by='composite_score', ascending=False
    ).reset_index(drop=True)

    # تصمیم
    always_enter_mask = (
        (result_df_filtered['composite_score'] == RISK_FREE_SCORE_SENTINEL)
        | (result_df_filtered['composite_score'] >= NEAR_EXPIRY_SCORE_BASE)
    )
    result_df_filtered['decision'] = np.where(
        always_enter_mask,
        'ENTER',
        np.where(
            result_df_filtered['composite_score'] > BCS_DECISION_THRESHOLD,
            'ENTER',
            'SKIP',
        ),
    )

    result_df_filtered['regime'] = np.select(
        [
            result_df_filtered['composite_score'] == RISK_FREE_SCORE_SENTINEL,
            result_df_filtered['composite_score'] >= NEAR_EXPIRY_SCORE_BASE,
        ],
        ['RISK_FREE', 'NEAR_EXPIRY'],
        default='NORMAL',
    )

    # ستون‌های نهایی (بدون اطلاعات اضافی، با حفظ محاسباتی‌ها)
    column_order = [
        'underlying', 'regime', 'decision', 'composite_score',
        'stock_price', 'long_option_symbol', 'long_strike',
        'long_ask_premium', 'short_option_symbol', 'short_strike',
        'short_bid_premium', 'capital_at_risk', 'max_net_profit',
        'max_profit_percent', 'monthly_return_%', 'margin30',
        'raw_return_percent', 'raw_margin_percent',
        'break_even_price', 'break_even_percent',
        'break_even_percent_scale', 'risk_reward_ratio',
        'days_to_maturity', 'long_volume', 'short_volume',
    ]
    column_order = [c for c in column_order if c in result_df_filtered.columns]
    result_df_filtered = result_df_filtered[column_order]

    # جایگزینی Sentinelها
    result_df_filtered['break_even_percent'] = result_df_filtered[
        'break_even_percent'
    ].replace(RISK_FREE_BREAK_EVEN_SENTINEL, 'Risk Free')
    result_df_filtered['break_even_percent_scale'] = result_df_filtered[
        'break_even_percent_scale'
    ].replace(RISK_FREE_BREAK_EVEN_SENTINEL, 'Risk Free')
    result_df_filtered['monthly_return_%'] = result_df_filtered[
        'monthly_return_%'
    ].replace(RISK_FREE_RETURN_SENTINEL, 'Infinite')
    result_df_filtered['margin30'] = result_df_filtered['margin30'].apply(
        lambda v: 'Risk Free' if v == np.inf else round(v, 2)
    )
    result_df_filtered['composite_score'] = result_df_filtered[
        'composite_score'
    ].replace(RISK_FREE_SCORE_SENTINEL, 'Risk Free')

    print(f"Bull Call Spread: {len(result_df_filtered)} positions")
    return result_df_filtered


# ==========================================================================================
# بخش ۸: اجرای Covered Call
# ==========================================================================================

def run_covered_call(df_options, df_vol):
    """اجرای کامل استراتژی Covered Call."""
    print("\n" + "=" * 65)
    print("Running Covered Call Strategy")
    print("=" * 65)

    results_fee = []

    for underlying_symbol, group in df_options.groupby('UnderlyingTicker'):
        market = get_symbol_market(underlying_symbol)
        kind = get_symbol_kind(underlying_symbol)

        opt_sell_commission = get_commission_rate(market, 'option', False)
        stock_buy_commission = get_commission_rate(market, kind, True)
        exercise_fee_rate = get_exercise_fee_rate(market, kind)
        exercise_tax_rate = EXERCISE_TAX_RATE

        for _, item in group.iterrows():
            ticker = item['Ticker']
            strike_price = item['StrikePrice']
            premium_call = item['BidPrice']
            stock_price = item['UnderlyingPrice']
            contract_size = item['ContractSize']
            days = item['DaysToMaturity']

            if premium_call <= 0 or stock_price <= 0:
                continue

            res = covered_call_analysis(
                stock_price=stock_price,
                strike_price=strike_price,
                premium_call=premium_call,
                contract_size=contract_size,
                opt_sell_commission=opt_sell_commission,
                stock_buy_commission=stock_buy_commission,
                exercise_fee_rate=exercise_fee_rate,
                exercise_tax_rate=exercise_tax_rate,
                days=days,
            )

            if res['status'] == 'DISCARD':
                continue

            if res['risk_reward_ratio'] < CC_MIN_RR_RATIO:
                continue

            results_fee.append({
                'underlying': underlying_symbol,
                'is_near_expiry': res['is_near_expiry'],
                'stock_price': round(stock_price, 0),
                'option_symbol': ticker,
                'strike': strike_price,
                'premium': round(premium_call, 0),
                'contract_size': contract_size,
                'capital_at_risk': res['capital_at_risk'],
                'net_profit': res['net_profit'],
                'profit_percent': res['profit_percent'],
                'monthly_return_%': res['monthly_return'],
                'raw_return_percent': res['raw_return_percent'],
                'raw_margin_percent': res['raw_margin_percent'],
                'break_even_price': res['break_even_price'],
                'break_even_percent': res['break_even_percent'],
                'break_even_percent_scale': res['break_even_percent_scale'],
                'max_drop_%': res['max_drop_percent'],
                'risk_reward_ratio': res['risk_reward_ratio'],
                'days_to_maturity': days,
                'volume': int(item.get('Volume', 0)),
            })

    if not results_fee:
        print("No valid Covered Call setups.")
        return pd.DataFrame()

    result_df = pd.DataFrame(results_fee)

    # Merge با نوسان
    if df_vol is not None and not df_vol.empty:
        result_df = pd.merge(result_df, df_vol, on='underlying', how='left')
    else:
        result_df['HV_60'] = np.nan
        result_df['VolatilityQualityScore'] = np.nan

    # فیلترینگ سخت
    result_df_filtered = apply_hard_constraints(
        result_df,
        min_monthly_return=CC_MIN_MONTHLY_RETURN,
        min_margin_floor=CC_MIN_MARGIN_FLOOR,
        margin_percentile=CC_MARGIN_PERCENTILE,
        min_m_risk=CC_MIN_M_RISK,
    )

    if result_df_filtered.empty:
        print("No valid setups after hard constraints.")
        return result_df_filtered

    # امتیازدهی
    result_df_filtered = calculate_composite_score(
        result_df_filtered,
        w_r=CC_W_R,
        w_m=CC_W_M,
        vol_quality_power=CC_VOL_QUALITY_POWER,
        imbalance_power=CC_IMBALANCE_POWER,
    )

    # رتبه‌بندی
    result_df_filtered = result_df_filtered.sort_values(
        by='composite_score', ascending=False
    ).reset_index(drop=True)

    # تصمیم
    always_enter_mask = (
        (result_df_filtered['composite_score'] == RISK_FREE_SCORE_SENTINEL)
        | (result_df_filtered['composite_score'] >= NEAR_EXPIRY_SCORE_BASE)
    )
    result_df_filtered['decision'] = np.where(
        always_enter_mask,
        'ENTER',
        np.where(
            result_df_filtered['composite_score'] > CC_DECISION_THRESHOLD,
            'ENTER',
            'SKIP',
        ),
    )

    result_df_filtered['regime'] = np.select(
        [
            result_df_filtered['composite_score'] == RISK_FREE_SCORE_SENTINEL,
            result_df_filtered['composite_score'] >= NEAR_EXPIRY_SCORE_BASE,
        ],
        ['RISK_FREE', 'NEAR_EXPIRY'],
        default='NORMAL',
    )

    # ستون‌های نهایی (بدون اطلاعات اضافی، با حفظ محاسباتی‌ها)
    column_order = [
        'underlying', 'regime', 'decision', 'composite_score',
        'stock_price', 'option_symbol', 'strike', 'premium',
        'contract_size', 'capital_at_risk', 'net_profit',
        'profit_percent', 'monthly_return_%', 'margin30',
        'raw_return_percent', 'raw_margin_percent',
        'break_even_price', 'break_even_percent',
        'break_even_percent_scale', 'max_drop_%',
        'risk_reward_ratio', 'days_to_maturity', 'volume',
    ]
    column_order = [c for c in column_order if c in result_df_filtered.columns]
    result_df_filtered = result_df_filtered[column_order]

    # جایگزینی Sentinelها
    result_df_filtered['break_even_percent'] = result_df_filtered[
        'break_even_percent'
    ].replace(RISK_FREE_BREAK_EVEN_SENTINEL, 'Risk Free')
    result_df_filtered['break_even_percent_scale'] = result_df_filtered[
        'break_even_percent_scale'
    ].replace(RISK_FREE_BREAK_EVEN_SENTINEL, 'Risk Free')
    result_df_filtered['monthly_return_%'] = result_df_filtered[
        'monthly_return_%'
    ].replace(RISK_FREE_RETURN_SENTINEL, 'Infinite')
    result_df_filtered['margin30'] = result_df_filtered['margin30'].apply(
        lambda v: 'Risk Free' if v == np.inf else round(v, 2)
    )
    result_df_filtered['composite_score'] = result_df_filtered[
        'composite_score'
    ].replace(RISK_FREE_SCORE_SENTINEL, 'Risk Free')

    print(f"Covered Call: {len(result_df_filtered)} positions")
    return result_df_filtered


# ==========================================================================================
# بخش ۹: ذخیره در اکسل (دو شیت)
# ==========================================================================================

def save_results_to_excel(bcs_df, cc_df, filename=OUTPUT_FILE_NAME):
    """ذخیره نتایج دو استراتژی در یک فایل اکسل با دو شیت."""
    filepath = current_dir / filename

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(
        start_color='203764', end_color='203764', fill_type='solid'
    )
    alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    body_font = Font(name='Segoe UI', size=10)
    gray_font = Font(color='808080', italic=True, name='Segoe UI', size=10)

    # rename dicts
    bcs_rename = {
        'underlying': 'نماد پایه',
        'regime': 'رژیم',
        'decision': 'تصمیم',
        'composite_score': 'امتیاز نهایی',
        'stock_price': 'قیمت سهم',
        'long_option_symbol': 'نماد خرید',
        'long_strike': 'قیمت اعمال خرید',
        'long_ask_premium': 'پریمیوم خرید',
        'short_option_symbol': 'نماد فروش',
        'short_strike': 'قیمت اعمال فروش',
        'short_bid_premium': 'پریمیوم فروش',
        'capital_at_risk': 'سرمایه درگیر',
        'max_net_profit': 'حداکثر سود خالص',
        'max_profit_percent': 'درصد بازدهی',
        'monthly_return_%': 'سود ماهانه (R30)',
        'margin30': 'حاشیه ماهانه (M30)',
        'raw_return_percent': 'سود خام',
        'raw_margin_percent': 'حاشیه خام',
        'break_even_price': 'قیمت سربه‌سر',
        'break_even_percent': 'درصد رشد تا سربه‌سر',
        'break_even_percent_scale': 'مقیاس سربه‌سر',
        'risk_reward_ratio': 'R/R',
        'days_to_maturity': 'روز تا سررسید',
        'long_volume': 'حجم خرید',
        'short_volume': 'حجم فروش',
    }

    cc_rename = {
        'underlying': 'نماد پایه',
        'regime': 'رژیم',
        'decision': 'تصمیم',
        'composite_score': 'امتیاز نهایی',
        'stock_price': 'قیمت سهم',
        'option_symbol': 'نماد اختیار',
        'strike': 'قیمت اعمال',
        'premium': 'پریمیوم فروش',
        'contract_size': 'اندازه قرارداد',
        'capital_at_risk': 'سرمایه درگیر',
        'net_profit': 'سود خالص',
        'profit_percent': 'درصد سود',
        'monthly_return_%': 'سود ماهانه (R30)',
        'margin30': 'حاشیه ماهانه (M30)',
        'raw_return_percent': 'سود خام',
        'raw_margin_percent': 'حاشیه خام',
        'break_even_price': 'قیمت سربه‌سر',
        'break_even_percent': 'درصد رشد تا سربه‌سر',
        'break_even_percent_scale': 'مقیاس سربه‌سر',
        'max_drop_%': 'حداکثر افت مجاز',
        'risk_reward_ratio': 'R/R',
        'days_to_maturity': 'روز تا سررسید',
        'volume': 'حجم',
    }

    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        # ===== شیت Bull Call Spread =====
        if bcs_df is not None and not bcs_df.empty:
            bcs_renamed = bcs_df.rename(columns=bcs_rename)
            bcs_renamed.to_excel(
                writer, sheet_name='bull_call_spread', index=False
            )
            ws = writer.sheets['bull_call_spread']
            _format_worksheet(ws, bcs_renamed, header_font, header_fill,
                              alignment, body_font, gray_font)
        else:
            # شیت خالی
            pd.DataFrame({'پیام': ['هیچ موقعیتی یافت نشد']}).to_excel(
                writer, sheet_name='bull_call_spread', index=False
            )

        # ===== شیت Covered Call =====
        if cc_df is not None and not cc_df.empty:
            cc_renamed = cc_df.rename(columns=cc_rename)
            cc_renamed.to_excel(
                writer, sheet_name='covered_call', index=False
            )
            ws = writer.sheets['covered_call']
            _format_worksheet(ws, cc_renamed, header_font, header_fill,
                              alignment, body_font, gray_font)
        else:
            pd.DataFrame({'پیام': ['هیچ موقعیتی یافت نشد']}).to_excel(
                writer, sheet_name='covered_call', index=False
            )

    print(f"\nResults saved to: {filename}")
    print(f"  - Sheet 'bull_call_spread': {len(bcs_df) if bcs_df is not None else 0} rows")
    print(f"  - Sheet 'covered_call': {len(cc_df) if cc_df is not None else 0} rows")


def _format_worksheet(ws, df_renamed, header_font, header_fill,
                       alignment, body_font, gray_font):
    """فرمت‌بندی یک worksheet."""
    # هدر
    for col_idx in range(1, len(df_renamed.columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment

    # بدنه
    for row_idx in range(2, len(df_renamed) + 2):
        for col_idx in range(1, len(df_renamed.columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is None or pd.isna(val) or val == "":
                cell.value = "-"
                cell.font = gray_font
            else:
                cell.font = body_font
            cell.alignment = alignment

    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(df_renamed.columns))}"
        f"{len(df_renamed) + 1}"
    )
    ws.freeze_panes = 'A2'

    # تنظیم عرض
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            if cell.value:
                text = str(cell.value)
                if '\n' in text:
                    lines = text.split('\n')
                    line_length = max(len(line) for line in lines)
                else:
                    line_length = len(text)
                if line_length > max_length:
                    max_length = line_length
        adjusted_width = min(max_length + 5, 50)
        ws.column_dimensions[column].width = adjusted_width


# ==========================================================================================
# بخش ۱۰: تابع اصلی
# ==========================================================================================

def main():
    """تابع اصلی اجرای استراتژی‌ها."""
    try:
        print("=" * 65)
        print("Strategies Analyzer (Bull Call Spread + Covered Call)")
        print("=" * 65)
        print(f"  BCS Enabled: {BCS_ENABLED}")
        print(f"  CC Enabled:  {CC_ENABLED}")
        print("=" * 65)

        if not BCS_ENABLED and not CC_ENABLED:
            print("No strategy is enabled. Exiting.")
            return

        # ===== بارگذاری پروفایل نوسان (مشترک) =====
        df_vol = load_volatility_profile()
        if df_vol.empty:
            print("\nWARNING: Volatility profile not loaded.")
        else:
            print(f"\nVolatility profile: {len(df_vol)} symbols loaded")

        # ===== بارگذاری داده بازار (مشترک) =====
        filtered_data = load_and_filter_data()

        if filtered_data.empty:
            print("No market data found.")
            return

        print(f"\nMarket data: {len(filtered_data)} CALL options loaded")

        # ===== اجرای Bull Call Spread =====
        bcs_results = None
        if BCS_ENABLED:
            bcs_results = run_bull_call_spread(filtered_data, df_vol)
        else:
            print("\n[SKIP] Bull Call Spread is disabled.")

        # ===== اجرای Covered Call =====
        cc_results = None
        if CC_ENABLED:
            cc_results = run_covered_call(filtered_data, df_vol)
        else:
            print("\n[SKIP] Covered Call is disabled.")

        # ===== ذخیره در اکسل =====
        save_results_to_excel(bcs_results, cc_results)

    except Exception as e:
        import traceback
        traceback.print_exc()


if __name__ == "__main__":
    main()