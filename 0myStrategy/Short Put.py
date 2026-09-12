# short_put_strategy.py
# -*- coding: utf-8 -*-

import sys
from pathlib import Path

# تنظیم مسیر پروژه
current_file_path = Path(__file__).resolve()
current_dir = current_file_path.parent
root_dir = current_dir.parent
sys.path.append(str(root_dir))

from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font, Alignment
import pandas as pd
from datetime import datetime
from data.downloader import MarketDownloader
from data.cleaner import DataCleaner
from config import (
    get_commission_rate,
    get_exercise_fee_rate,
    get_symbol_kind,
    get_symbol_market,
)


def short_put_with_fees(premium_put, stock_price, strike_price, contract_size,
                        opt_sell_commission, exercise_fee_rate, days,
                        margin_rate=0.20):
    """
    محاسبه بازده استراتژی Short Put با احتساب کارمزدها و وجه تضمین

    منطق Short Put:
    - فروشنده Put متعهد می‌شود در صورت اعمال، دارایی را به قیمت strike بخرد.
    - در ازای این تعهد، پریمیوم دریافت می‌کند.
    - سود حداکثر = پریمیوم دریافتی - کارمزدها (اگر Put بی‌ارزش منقضی شود)
    - حداکثر زیان = (strike - 0) * size - premium  (وقتی قیمت پایه به صفر برسد)
    - وجه تضمین (margin) توسط بورس دریافت می‌شود که در محاسبه بازده لحاظ می‌شود.
    """

    # ========== 1. دریافت پریمیوم (ورود نقدی) ==========
    premium_total = round(premium_put * contract_size, 0)
    entry_fee = -round(premium_total * opt_sell_commission, 0)

    # ========== 2. وجه تضمین (Margin) ==========
    # تقریبی: درصدی از (قیمت اعمال × اندازه قرارداد)
    # برای محاسبه بازده واقعی باید در مخرج کسر بیاید
    margin_required = round(strike_price * contract_size * margin_rate, 0)

    # ========== 3. جریان نقدی اولیه (سرمایه درگیر) ==========
    # وجه تضمین به عنوان سرمایه درگیر محسوب می‌شود
    initial_cash_flow = premium_total + entry_fee  # سود اولیه از پریمیوم
    capital_at_risk = margin_required  # سرمایه درگیر

    # ========== 4. ارزش ذاتی در سررسید (بدهی فروشنده) ==========
    # اگر stock < strike، فروشنده متضرر می‌شود (باید به strike بخرد)
    intrinsic_liability = max(0, strike_price - stock_price) * contract_size

    # ========== 5. کارمزد اعمال (فقط اگر ITM باشد و اعمال شود) ==========
    exercise_fee = 0
    if stock_price < strike_price:
        # در اعمال Put، فروشنده دارایی را به قیمت strike می‌خرد
        settlement_amount = strike_price * contract_size
        exercise_fee = -round(settlement_amount * exercise_fee_rate, 0)

    # ========== 6. سود خالص نهایی ==========
    # net_profit = پریمیوم دریافتی - کارمزدها - بدهی ارزش ذاتی
    net_profit = initial_cash_flow - intrinsic_liability + exercise_fee

    # ========== 7. بازده درصدی نسبت به سرمایه درگیر (وجه تضمین) ==========
    profit_percent = round(
        (net_profit / capital_at_risk) * 100, 2
    ) if capital_at_risk != 0 else 0
    monthly_return = round(profit_percent * (30 / days), 2)

    # ========== 8. نقطه سربه‌سر ==========
    # در Short Put: break_even = strike - (premium - fees per share)
    # چون فروشنده پریمیوم می‌گیرد، سربه‌سر پایین‌تر از strike می‌رود
    total_credit_per_share = premium_put - (abs(entry_fee) / contract_size)
    if stock_price < strike_price:
        total_credit_per_share -= abs(exercise_fee) / contract_size
    break_even_price = round(strike_price - total_credit_per_share, 0)

    # ========== 9. درصد فاصله قیمت پایه فعلی تا نقطه سربه‌سر ==========
    # در Short Put: اگر stock > break_even باشد، فاصله امن داریم
    # (قیمت پایه باید بیفتد تا به سربه‌سر برسد)
    if break_even_price != 0:
        break_even_percent = round(
            ((stock_price - break_even_price) / stock_price) * 100, 2
        )
    else:
        break_even_percent = 0

    # ========== 10. درصد فاصله قیمت پایه تا strike ==========
    # در Short Put: OTM مطلوب است (stock > strike)
    # اگر stock < strike باشیم یعنی ITM و ریسک اعمال بالا
    if stock_price > 0 and strike_price > 0:
        distance_to_strike = ((stock_price - strike_price) / stock_price) * 100
        # در Short Put: مثبت بودن یعنی OTM (امن)
        max_loss_price_percent = round(distance_to_strike, 2) if distance_to_strike > 0 else 0.0
    else:
        max_loss_price_percent = 0.0

    # ========== 11. درصد ریسک نسبت به حاشیه امنیت ==========
    # در Short Put: هرچه OTM‌تر باشیم ریسک کمتر (پرتوقع‌تر)
    if stock_price > strike_price:
        price_difference = stock_price - strike_price
        if price_difference > 0:
            total_premium_income = premium_put * contract_size
            risk_percent = round(
                (total_premium_income / (price_difference * contract_size)) * 100, 2
            )
        else:
            risk_percent = 100.0
    else:
        # ITM => ریسک بالا
        risk_percent = 100.0

    return {
        'net_profit': net_profit,
        'profit_percent': profit_percent,
        'monthly_return': monthly_return,
        'break_even_price': break_even_price,
        'break_even_percent': break_even_percent,
        'intrinsic_liability': intrinsic_liability,
        'fees_total': entry_fee + exercise_fee,
        'margin_required': margin_required,
        'max_loss_price_percent': max_loss_price_percent,
        'risk_percent': risk_percent,
    }


def load_and_filter_data():
    """
    بارگذاری داده‌ها از TSE و فیلتر کردن گزینه‌های اختیار فروش (PUT)
    برای استراتژی Short Put، BidPrice مهم است (چون می‌خواهیم بفروشیم)
    """
    df_raw = MarketDownloader.from_tsetmc_direct()
    df_cleaned = DataCleaner.clean(df_raw)
    df_final = DataCleaner.add_derived_columns(df_cleaned)

    # ------ تغییر دستی قیمت نماد پایه (اختیاری) ------
    # UnderlyingTicker = 'خودرو'
    # if UnderlyingTicker in df_final['UnderlyingTicker'].values:
    #     mask_self = df_final['UnderlyingTicker'] == UnderlyingTicker
    #     df_final.loc[mask_self, 'UnderlyingPrice'] = 528

    # فیلتر گزینه‌های اختیار فروش (PUT)
    filter_option = df_final[
        (df_final['DaysToMaturity'] > 2.0) &
        (df_final['Type'].apply(lambda x: x.name == 'PUT'))
    ].copy()

    # حذف موارد نامطلوب
    EXCLUDED_UNDERLYING = ['اهرم']
    EXCLUDED_NAME_PATTERN = ['1405/04', '1405-04']
    exclude_mask = (
        (filter_option['UnderlyingTicker'].isin(EXCLUDED_UNDERLYING)) &
        (filter_option['Name'].str.contains(
            '|'.join(EXCLUDED_NAME_PATTERN), na=False))
    )
    filter_option = filter_option[~exclude_mask].copy()

    return filter_option


def run_short_put_strategy(df_options, max_break_even_percent=12):
    """
    اجرای استراتژی Short Put روی داده‌های فیلتر شده

    برای Short Put، از BidPrice استفاده می‌کنیم چون می‌خواهیم بفروشیم.
    اگر BidPrice موجود نبود، از AskPrice به عنوان تخمین استفاده می‌شود.
    """
    results_fee = []

    for underlying_symbol, group in df_options.groupby('UnderlyingTicker'):
        market = get_symbol_market(underlying_symbol)
        kind = get_symbol_kind(underlying_symbol)

        # کارمزد فروش اختیار (is_buy=False)
        opt_sell_commission = get_commission_rate(market, 'option', False)
        exercise_fee_rate = get_exercise_fee_rate(market, kind)

        for _, item in group.iterrows():
            ticker = item['Ticker']
            strike_price = item['StrikePrice']

            # برای Short Put از BidPrice استفاده می‌کنیم (قیمتی که می‌توانیم بفروشیم)
            bid_price = item.get('BidPrice', None)
            ask_price = item.get('AskPrice', None)
            if bid_price is not None and not pd.isna(bid_price) and bid_price > 0:
                premium_put = bid_price
            elif ask_price is not None and not pd.isna(ask_price):
                premium_put = ask_price  # تخمین
            else:
                continue  # بدون قیمت قابل معامله، رد می‌شود

            stock_price = item['UnderlyingPrice']
            contract_size = item['ContractSize']
            days = item['DaysToMaturity']

            # محاسبات با کارمزد
            results = short_put_with_fees(
                premium_put, stock_price, strike_price, contract_size,
                opt_sell_commission, exercise_fee_rate, days)

            max_loss_price_percent_scale = results['max_loss_price_percent'] * ((30 / days) ** 0.5)
            break_even_percent_scale = results['break_even_percent'] * ((30 / days) ** 0.5)

            results_fee.append({
                'underlying': underlying_symbol,
                'stock_price': round(stock_price, 0),
                'option_symbol': ticker,
                'strike': strike_price,
                'premium': round(premium_put, 0),
                'net_profit': results['net_profit'],
                'profit_percent': results['profit_percent'],
                'monthly_return_%': results['monthly_return'],
                'break_even_price': results['break_even_price'],
                'break_even_percent': results['break_even_percent'],
                'break_even_percent_scale': round(break_even_percent_scale, 2),
                'max_loss_price_percent': results['max_loss_price_percent'],
                'max_loss_price_percent_scale': round(max_loss_price_percent_scale, 2),
                'margin_required': results['margin_required'],
                'risk_percent': results['risk_percent'],
                'days_to_maturity': days,
                'volume': int(item.get('Volume', 0)),
            })

    result_df = pd.DataFrame(results_fee)

    if result_df.empty:
        return result_df

    # ------ فیلتر و مرتب‌سازی مخصوص Short Put ------
    # در Short Put، هدف: پریمیوم بالا، OTM بودن، فاصله امن زیاد
    # فیلتر: حذف Put‌هایی که خیلی ضررده هستند (break_even_percent خیلی منفی)
    result_df_filtered = result_df[
        result_df['break_even_percent'] >= -max_break_even_percent
    ].copy()

    # مرتب‌سازی: نزولی بر اساس بازده ماهانه (بهترین = بیشترین درآمد ماهانه)
    result_df_filtered = result_df_filtered.sort_values(
        ['monthly_return_%', 'break_even_percent'],
        ascending=[False, False]
    ).reset_index(drop=True)

    return result_df_filtered


def save_results_to_excel(result_df, filename="result_short_put.xlsx"):
    """
    ذخیره نتایج در فایل اکسل با استایل‌بندی حرفه‌ای
    """
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='7B2E2E',
                              end_color='7B2E2E', fill_type='solid')
    alignment = Alignment(horizontal='center',
                          vertical='center', wrap_text=True)
    body_font = Font(name='Segoe UI', size=10)
    gray_font = Font(color='808080', italic=True, name='Segoe UI', size=10)

    result_df = result_df.rename(columns={
        'stock_price': 'قیمت نماد پایه',
        'strike': 'قیمت اعمال',
        'premium': 'پریمیوم دریافتی (Bid)',
        'monthly_return_%': 'درصد سود ماهانه\n(نسبت به وجه تضمین)',
        'break_even_price': 'قیمت سربه‌سر',
        'break_even_percent': 'درصد فاصله تا نقطه سربه‌سر\n(هرچه بیشتر = امن‌تر)',
        'break_even_percent_scale': 'مقیاس درصد فاصله تا سربه‌سر\n(به نسبت 30 روز)',
        'max_loss_price_percent': 'درصد OTM بودن\n(هرچه بیشتر = امن‌تر)',
        'max_loss_price_percent_scale': 'مقیاس OTM\n(به نسبت 30 روز)',
        'margin_required': 'وجه تضمین تقریبی',
        'volume': 'حجم معاملات روز جاری',
        'risk_percent': 'درصد ریسک\n(هرچه کمتر = بهتر)',
    })

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename_with_time = f"result_short_put_{timestamp}.xlsx"
    filename_with_time = 'result_short_put.xlsx'
    filepath = Path(__file__).parent / filename_with_time

    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        result_df.to_excel(writer, sheet_name='short_put', index=False)
        worksheet = writer.sheets['short_put']

        # استایل هدر
        for col_idx in range(1, len(result_df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = alignment

        # استایل بدنه
        columns_list = result_df.columns.tolist()
        for row_idx, row in enumerate(result_df.itertuples(index=False), start=2):
            for col_idx, col_name in enumerate(columns_list, start=1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                val = row[col_idx - 1]
                if val is None or pd.isna(val):
                    cell.value = "-"
                    cell.font = gray_font
                else:
                    cell.font = body_font
                cell.alignment = alignment

        # هایلایت ماکزیمم/مینیمم
        highlight_columns = ['monthly_return_%', 'break_even_percent', 'profit_percent']
        max_fill = PatternFill(start_color='92D050', end_color='92D050', fill_type='solid')
        min_fill = PatternFill(start_color='FF9999', end_color='FF9999', fill_type='solid')

        col_indices = {}
        for col_idx, col_name in enumerate(columns_list, start=1):
            if col_name in highlight_columns:
                col_indices[col_name] = col_idx

        for col_name, col_idx in col_indices.items():
            values = []
            for row_idx in range(2, len(result_df) + 2):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                if cell.value is not None and cell.value != "-":
                    try:
                        values.append(float(cell.value))
                    except Exception:
                        pass

            if values:
                max_val = max(values)
                min_val = min(values)
                for row_idx in range(2, len(result_df) + 2):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    if cell.value is not None and cell.value != "-":
                        try:
                            val = float(cell.value)
                            if val == max_val:
                                cell.fill = max_fill
                            elif val == min_val:
                                cell.fill = min_fill
                        except Exception:
                            pass

        # فیلتر و Freeze
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(result_df.columns))}{len(result_df) + 1}"
        worksheet.freeze_panes = 'A2'

        # عرض ستون‌ها
        for col in worksheet.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                if cell.value:
                    text = str(cell.value)
                    if '\n' in text:
                        lines = text.split('\n')
                        line_length = max(len(line) for line in lines)
                    else:
                        line_length = len(text)
                    if line_length > max_length:
                        max_length = line_length
            adjusted_width = min(max_length + 5, 50)
            worksheet.column_dimensions[column].width = adjusted_width

    print(f"result {filename_with_time}  save")
    return filename_with_time


def main():
    """
    تابع اصلی اجرای استراتژی Short Put
    """
    try:
        # 1. بارگذاری و فیلتر داده‌ها
        filtered_data = load_and_filter_data()

        if filtered_data.empty:
            return

        # 2. اجرای استراتژی
        results = run_short_put_strategy(
            filtered_data, max_break_even_percent=12)

        if results.empty:
            return

        # 3. ذخیره نتایج
        save_results_to_excel(results)

    except Exception as e:
        import traceback
        traceback.print_exc()


if __name__ == "__main__":
    main()