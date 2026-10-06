# 0myStrategy/bcs_scanner.py
# -*- coding: utf-8 -*-

"""
bcs_scanner.py — اسکن BCS با تزریق ریزمعاملات در df_final
============================================================

نحوه اجرا:
    python bcs_scanner.py

پیش‌نیاز:
    - fetch_option_inscodes.py  (تولید option_inscodes.json)
    - fetch_trade_history.py    (تولید trades/YYYYMMDD/)
"""

import sys
import json
from pathlib import Path
from datetime import datetime, time as dtime

import numpy as np
import pandas as pd

# --- مسیرها ---
CURRENT_DIR = Path(__file__).resolve().parent
ROOT_DIR = CURRENT_DIR.parent
sys.path.append(str(ROOT_DIR))
sys.path.append(str(CURRENT_DIR))

TRADES_DIR = CURRENT_DIR / "trades"
OPTIONS_JSON = CURRENT_DIR / "option_inscodes.json"

# --- تنظیمات ---
DATE_STR = None
MARKET_CLOSE = dtime(12, 30, 0)


# --- Import ---
from Bull_call_spread import (
    run_bull_call_spread_strategy,
    load_volatility_profile,
)

from data.cleaner import DataCleaner
from data.downloader import MarketDownloader


# =====================================================================
# بارگذاری registry از JSON
# =====================================================================

def load_registry_from_json():
    """بارگذاری OPTIONS_REGISTRY از option_inscodes.json."""
    with open(OPTIONS_JSON, 'r', encoding='utf-8') as f:
        data = json.load(f)
    return data['registry']


# =====================================================================
# بارگذاری ریزمعاملات
# =====================================================================

def load_trades_from_cache(date_str, registry):
    """بارگذاری ریزمعاملات همه نمادهای registry."""
    folder = TRADES_DIR / date_str
    if not folder.exists():
        raise FileNotFoundError(f"Trades folder not found: {folder}")

    # --- نگاشت insCode -> ticker ---
    code_to_ticker = {}
    for underlying, options in registry.items():
        for ticker, code in options.items():
            code_to_ticker[code] = ticker

    trades_by_ticker = {}
    for parquet_file in folder.glob("trades_*.parquet"):
        try:
            ins_code = int(parquet_file.stem.replace("trades_", ""))
        except ValueError:
            continue
        ticker = code_to_ticker.get(ins_code)
        if ticker is None:
            continue
        try:
            df = pd.read_parquet(parquet_file)
            if not df.empty:
                trades_by_ticker[ticker] = df
                print(f"  [OK]   {ticker}: {len(df)} trades")
        except Exception as e:
            print(f"  [WARN] {parquet_file.name}: {e}")

    return trades_by_ticker


# =====================================================================
# توابع کمکی
# =====================================================================

def group_long_trades_by_time(long_df):
    """گروه‌بندی معاملات بر اساس time (max قیمت + تعداد)."""
    return long_df.groupby('time').agg(
        max_price=('price', 'max'),
        count=('price', 'size'),
    ).reset_index().sort_values('time').reset_index(drop=True)


def find_first_short_after(short_df, T_L, T_L_next):
    """اولین معامله short در پنجره (T_L, T_L_next)."""
    candidates = short_df[short_df['time'] > T_L]
    if candidates.empty:
        return None
    first = candidates.iloc[0]
    T_S = first['time']
    if T_S >= T_L_next:
        return None
    return (T_S, first['price'], first['nTran'], first['volume'])


# =====================================================================
# هسته اسکن — تزریق و پاس دادن
# =====================================================================

def scan_all(df_ours, trades_by_ticker, df_vol):
    """
    برای هر جفت و هر لحظه T_L:
      1. کپی از df_ours
      2. تزریق قیمت خرید (long) و فروش (short)
      3. فیلتر فقط نمادهای آن جفت
      4. پاس دادن به run_bull_call_spread_strategy
    """
    all_results = []

    for (underlying, days), group in df_ours.groupby(['UnderlyingTicker', 'DaysToMaturity']):
        group_sorted = group.sort_values('StrikePrice')
        options_list = group_sorted.to_dict('records')
        n = len(options_list)

        if n < 2:
            continue

        print(f"\n  -> {underlying} | {days} days | {n} options")

        # --- همه جفت‌ها ---
        for i in range(n):
            for j in range(i + 1, n):
                L = options_list[i]   # strike پایین‌تر
                S = options_list[j]   # strike بالاتر

                long_ticker = L['Ticker']
                short_ticker = S['Ticker']

                long_df = trades_by_ticker.get(long_ticker)
                short_df = trades_by_ticker.get(short_ticker)

                if long_df is None or short_df is None:
                    continue
                if long_df.empty or short_df.empty:
                    continue

                # --- گروه‌بندی ریزمعاملات long ---
                long_groups = group_long_trades_by_time(long_df)
                if len(long_groups) == 0:
                    continue

                pair_tickers = [long_ticker, short_ticker]
                n_moments = 0

                # --- حلقه روی T_Lها ---
                for k in range(len(long_groups)):
                    g = long_groups.iloc[k]
                    T_L = g['time']
                    P_L = g['max_price']

                    # --- پنجره ---
                    T_L_next = (
                        long_groups.iloc[k + 1]['time']
                        if k + 1 < len(long_groups)
                        else MARKET_CLOSE
                    )

                    # --- اولین short ---
                    found = find_first_short_after(short_df, T_L, T_L_next)
                    if found is None:
                        continue

                    T_S, P_S, nTran_S, volume_S = found

                    # --- کپی از df_ours ---
                    df_copy = df_ours.copy()

                    # --- تزریق قیمت‌ها ---
                    df_copy.loc[df_copy['Ticker'] == long_ticker, 'AskPrice'] = P_L
                    df_copy.loc[df_copy['Ticker'] == short_ticker, 'BidPrice'] = P_S

                    # --- فقط نمادهای این جفت ---
                    df_pair = df_copy[df_copy['Ticker'].isin(pair_tickers)].copy()

                    # --- صفر کردن ستون‌های مخالف ---
                    df_pair.loc[df_pair['Ticker'] == long_ticker, 'BidPrice'] = 0
                    df_pair.loc[df_pair['Ticker'] == short_ticker, 'AskPrice'] = 0

                    # --- فراخوانی با پیش‌فرض‌های ماژول ---
                    try:
                        result_df = run_bull_call_spread_strategy(
                            df_pair, df_vol=df_vol
                        )
                    except Exception as e:
                        print(f"     [WARN] {long_ticker}/{short_ticker} @ {T_L}: {e}")
                        continue

                    if result_df is None or result_df.empty:
                        continue

                    # --- اضافه کردن اطلاعات زمانی و حجم ---
                    for _, row in result_df.iterrows():
                        rec = row.to_dict()
                        rec['time_L'] = T_L
                        rec['time_S'] = T_S
                        rec['nTran_S'] = nTran_S
                        rec['long_count'] = g['count']
                        rec['long_volume'] = int(g['count'])
                        rec['short_volume'] = int(volume_S)
                        all_results.append(rec)

                    n_moments += 1

                if n_moments > 0:
                    print(f"     [OK] {long_ticker}/{short_ticker}: {n_moments} moments")

    return pd.DataFrame(all_results)


# =====================================================================
# ذخیره اکسل
# =====================================================================

def save_to_excel(results_df, date_str):
    """ذخیره نتایج در اکسل با ستون‌های درخواستی."""
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    filepath = CURRENT_DIR / f"bcs_opportunities_{date_str}.xlsx"

    # --- rename به فارسی ---
    rename_dict = {
        'underlying': 'نماد پایه',
        'regime': 'رژیم',
        'decision': 'تصمیم',
        'composite_score': 'امتیاز نهایی',
        'HV_60': 'نوسان ۶۰ روزه',
        'stock_price': 'قیمت سهم',
        'long_option_symbol': 'نماد خرید',
        'time_L': 'زمان خرید',
        'long_strike': 'قیمت اعمال خرید',
        'long_ask_premium': 'پریمیوم خرید',
        'short_option_symbol': 'نماد فروش',
        'time_S': 'زمان فروش',
        'short_strike': 'قیمت اعمال فروش',
        'short_bid_premium': 'پریمیوم فروش',
        'capital_at_risk': 'سرمایه درگیر',
        'max_net_profit': 'حداکثر سود خالص',
        'max_profit_percent': 'درصد بازدهی',
        'monthly_return_%': 'سود ماهانه (R30)',
        'margin30': 'حاشیه ماهانه (M30)',
        'raw_return_percent': 'سود خام',
        'raw_margin_percent': 'حاشیه خام',
        'break_even_price': 'قیمت سربه‌سر',
        'break_even_percent': 'درصد رشد تا سربه‌سر',
        'break_even_percent_scale': 'مقیاس سربه‌سر',
        'risk_reward_ratio': 'R/R',
        'days_to_maturity': 'روز تا سررسید',
        'long_volume': 'حجم خرید',
        'short_volume': 'حجم فروش',
        'maturity_date': 'تاریخ اعمال',
    }

    # --- ترتیب دقیق ستون‌ها ---
    column_order = [
        'underlying', 'regime', 'decision', 'composite_score',
        'HV_60', 'stock_price',
        'long_option_symbol', 'time_L', 'long_strike', 'long_ask_premium',
        'short_option_symbol', 'time_S', 'short_strike', 'short_bid_premium',
        'capital_at_risk', 'max_net_profit', 'max_profit_percent',
        'monthly_return_%', 'margin30', 'raw_return_percent', 'raw_margin_percent',
        'break_even_price', 'break_even_percent', 'break_even_percent_scale',
        'risk_reward_ratio', 'days_to_maturity',
        'long_volume', 'short_volume', 'maturity_date',
    ]

    cols = [c for c in column_order if c in results_df.columns]
    df_out = results_df[cols].rename(columns=rename_dict)

    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        df_out.to_excel(writer, sheet_name='BCS Opportunities', index=False)
        ws = writer.sheets['BCS Opportunities']

        # --- استایل هدر ---
        header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(
            start_color='203764', end_color='203764', fill_type='solid'
        )
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)

        for col_idx in range(1, len(df_out.columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        # --- استایل بدنه ---
        body_font = Font(name='Segoe UI', size=10)
        for row_idx in range(2, len(df_out) + 2):
            for col_idx in range(1, len(df_out.columns) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = body_font
                cell.alignment = center

        # --- فیلتر و freeze ---
        ws.auto_filter.ref = (
            f"A1:{get_column_letter(len(df_out.columns))}{len(df_out) + 1}"
        )
        ws.freeze_panes = 'A2'

        # --- عرض ستون‌ها ---
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 5, 50)

    print(f"\n[EXCEL] Saved: {filepath}")
    return filepath


# =====================================================================
# تابع اصلی
# =====================================================================

def main():
    date_str = DATE_STR or datetime.now().strftime('%Y%m%d')

    print("=" * 70)
    print(f"BCS Scanner — {date_str}")
    print("=" * 70)

    # --- ۱. registry ---
    print("\n[STEP 1] Loading registry...")
    registry = load_registry_from_json()
    our_tickers = set()
    for underlying, options in registry.items():
        our_tickers.update(options.keys())
    print(f"  [OK] {len(our_tickers)} tickers")

    # --- ۲. ریزمعاملات ---
    print("\n[STEP 2] Loading trades...")
    try:
        trades_by_ticker = load_trades_from_cache(date_str, registry)
    except FileNotFoundError as e:
        print(f"\n[FAIL] {e}")
        return
    print(f"  [OK] {len(trades_by_ticker)} symbols")

    # --- ۳. df_final ---
    print("\n[STEP 3] Loading df_final...")
    df_raw = MarketDownloader.from_tsetmc_direct()
    df_cleaned = DataCleaner.clean(df_raw)
    df_final = DataCleaner.add_derived_columns(df_cleaned)

    df_ours = df_final[df_final['Ticker'].isin(our_tickers)].copy()
    print(f"  [OK] {len(df_ours)} options filtered")

    if df_ours.empty:
        print("[FAIL] No matching options")
        return

    # --- ۴. پروفایل نوسان ---
    print("\n[STEP 4] Loading volatility profile...")
    df_vol = load_volatility_profile()
    if df_vol.empty:
        print("  [WARN] Volatility profile not loaded")
    else:
        print(f"  [OK] {len(df_vol)} symbols")

    # --- ۵. اسکن ---
    print("\n[STEP 5] Scanning...")
    results_df = scan_all(df_ours, trades_by_ticker, df_vol)

    if results_df.empty:
        print("\n[WARN] No results found.")
        return

    print(f"\n[OK] Total: {len(results_df)} results")

    # --- ۶. فیلتر ENTER ---
    if 'decision' in results_df.columns:
        results_df = results_df[
            results_df['decision'] == 'ENTER'
        ].reset_index(drop=True)
        print(f"[FILTER] {len(results_df)} ENTER positions")

    if results_df.empty:
        print("[WARN] No ENTER positions.")
        return

    # --- ۷. ذخیره ---
    save_to_excel(results_df, date_str)


if __name__ == "__main__":
    main()